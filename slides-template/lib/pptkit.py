"""
pptkit — design system + low-level helpers for the "Aurora" report template.

Everything is native PowerPoint: shapes, text, pictures, charts, tables.
Morph transitions and entrance animations are written straight into the slide XML.
Units: inches (floats) unless noted.
"""
import copy
import os

from lxml import etree
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION, XL_MARKER_STYLE, XL_TICK_LABEL_POSITION
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt, Emu

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)

SW, SH = 13.333, 7.5          # 16:9 widescreen canvas
ML = 0.65                     # left / right margin

# ---------------------------------------------------------------- fonts
F_BLACK = 'Montserrat Black'
F_HEAD = 'Montserrat ExtraBold'
F_BOLD = 'Montserrat'          # used with b=1
F_SEMI = 'Montserrat SemiBold'
F_MED = 'Montserrat Medium'
F_BODY = 'Montserrat'

# ---------------------------------------------------------------- themes
THEMES = {
    # new signature palette: midnight navy + electric aqua glow
    'aurora': dict(
        name='Aurora',
        bg_light='F4F7FE', bg_dark='060E2A',
        ink='0B1A48', text='3B4866', muted='7F8BA6', soft='E6ECFA',
        card='0C1B49', card2='1B3FC4', accent='3DE3F0', accent_ink='06123A',
        violet='7A5CFF', hot='FF5D5D', up='12B886', line='C9D3EA',
        chart=['0C1B49', '2A55E5', '2BCBDD', '7A5CFF'],
        blob_light=[('3DE3F0', 62), ('7A8CFF', 40), ('9FE9FF', 55)],
        blob_dark=[('3DE3F0', 55), ('2A55E5', 70), ('7A5CFF', 45)],
        tint='0B1A48', grid_dark='27407A', row_alt='13265E',
    ),
    # homage palette: deep emerald + electric lime (closest to the reference deck)
    'lime': dict(
        name='Lime',
        bg_light='F6FAEE', bg_dark='071A17',
        ink='0C3A35', text='3B4E4A', muted='7E918C', soft='E8F1DA',
        card='0B2C28', card2='1D6B67', accent='C6F23A', accent_ink='0B2C28',
        violet='5FD3A6', hot='E5533D', up='4CAF50', line='CFDCC4',
        chart=['0B2C28', '1D6B67', '9ED82C', '5FD3A6'],
        blob_light=[('C6F23A', 70), ('9EE06A', 45), ('DDF79A', 60)],
        blob_dark=[('C6F23A', 55), ('1D6B67', 70), ('9ED82C', 45)],
        tint='0B2C28', grid_dark='24514B', row_alt='0F3A35',
    ),
}

T = dict(THEMES['aurora'])    # active theme (mutated by set_theme)


def set_theme(key):
    T.clear()
    T.update(THEMES[key])
    T['key'] = key


def rgb(h):
    return RGBColor.from_string(h)


def E(v):
    return Emu(int(round(v * 914400)))


def mix(c1, c2, t):
    """Blend hex colour c1 over c2 with opacity t (0-1) -> solid hex (renders identically everywhere)."""
    a = [int(c1[i:i + 2], 16) for i in (0, 2, 4)]
    b = [int(c2[i:i + 2], 16) for i in (0, 2, 4)]
    return ''.join(f'{round(x * t + y * (1 - t)):02X}' for x, y in zip(a, b))


FONT_FILES = {
    'Montserrat Black': 'Montserrat-Black.ttf', 'Montserrat ExtraBold': 'Montserrat-ExtraBold.ttf',
    'Montserrat SemiBold': 'Montserrat-SemiBold.ttf', 'Montserrat Medium': 'Montserrat-Medium.ttf',
    'Montserrat': 'Montserrat-Regular.ttf',
}
_font_cache = {}


def text_w(t, font, size):
    """Rendered width of a single line in inches (uses the bundled TTFs)."""
    from PIL import ImageFont
    key = (font, size)
    if key not in _font_cache:
        _font_cache[key] = ImageFont.truetype(os.path.join(ROOT, 'fonts', FONT_FILES.get(font, 'Montserrat-Regular.ttf')),
                                              int(size * 20))
    return _font_cache[key].getlength(t) / 20 / 72


def n_lines(t, font, size, w):
    total = 0
    for para in t.split('\n'):
        words, cur, lines = para.split(' '), '', 1
        for wd in words:
            trial = (cur + ' ' + wd).strip()
            if text_w(trial, font, size) <= w or not cur:
                cur = trial
            else:
                lines += 1
                cur = wd
        total += lines
    return total


def fit_size(t, font, size, w, max_lines=1, min_ratio=0.7):
    sz = size
    while sz > size * min_ratio and n_lines(t, font, sz, w * 0.97) > max_lines:
        sz -= 1
    return sz


NS = {
    'a': 'http://schemas.openxmlformats.org/drawingml/2006/main',
    'p': 'http://schemas.openxmlformats.org/presentationml/2006/main',
    'r': 'http://schemas.openxmlformats.org/officeDocument/2006/relationships',
    'c': 'http://schemas.openxmlformats.org/drawingml/2006/chart',
}


def X(xml):
    """Parse an XML fragment that uses a:/p:/r: prefixes."""
    decl = ' '.join(f'xmlns:{k}="{v}"' for k, v in NS.items())
    wrapped = f'<root {decl}>{xml}</root>'
    return etree.fromstring(wrapped)[0]


# ================================================================ presentation
def new_presentation():
    prs = Presentation()
    prs.slide_width = E(SW)
    prs.slide_height = E(SH)
    return prs


def blank_slide(prs, dark=False):
    layout = prs.slide_layouts[6]
    s = prs.slides.add_slide(layout)
    fill = s.background.fill
    fill.solid()
    fill.fore_color.rgb = rgb(T['bg_dark'] if dark else T['bg_light'])
    s._dark = dark
    s._anims = []
    s._morph = None
    return s


# ================================================================ low-level shape styling
def _spPr(shape):
    return shape._element.spPr


def set_fill(shape, color=None, alpha=None):
    if color is None:
        shape.fill.background()
        return
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(color)
    if alpha is not None:
        clr = _spPr(shape).find(qn('a:solidFill'))[0]
        a = etree.SubElement(clr, qn('a:alpha'))
        a.set('val', str(int(alpha * 1000)))


def set_line(shape, color=None, width=0.75, alpha=None, dash=None):
    ln = shape.line
    if color is None:
        ln.fill.background()
        return
    ln.color.rgb = rgb(color)
    ln.width = Pt(width)
    if alpha is not None:
        clr = _spPr(shape).find(qn('a:ln')).find(qn('a:solidFill'))[0]
        a = etree.SubElement(clr, qn('a:alpha'))
        a.set('val', str(int(alpha * 1000)))
    if dash:
        lnel = _spPr(shape).find(qn('a:ln'))
        d = etree.SubElement(lnel, qn('a:prstDash'))
        d.set('val', dash)


def _effect_lst(spPr):
    eff = spPr.find(qn('a:effectLst'))
    if eff is None:
        eff = etree.SubElement(spPr, qn('a:effectLst'))
        # effectLst must come after fill/ln
        ln = spPr.find(qn('a:ln'))
        if ln is not None:
            ln.addnext(eff)
    return eff


def shadow(shape, blur=0.25, dist=0.08, alpha=18, color='0B1A48', angle=90):
    spPr = _spPr(shape)
    eff = _effect_lst(spPr)
    sh = etree.SubElement(eff, qn('a:outerShdw'))
    sh.set('blurRad', str(int(blur * 914400)))
    sh.set('dist', str(int(dist * 914400)))
    sh.set('dir', str(int(angle * 60000)))
    sh.set('algn', 'ctr')
    sh.set('rotWithShape', '0')
    c = etree.SubElement(sh, qn('a:srgbClr'))
    c.set('val', color)
    a = etree.SubElement(c, qn('a:alpha'))
    a.set('val', str(int(alpha * 1000)))


def glow(shape, rad=0.25, color='3DE3F0', alpha=40):
    eff = _effect_lst(_spPr(shape))
    g = etree.SubElement(eff, qn('a:glow'))
    g.set('rad', str(int(rad * 914400)))
    c = etree.SubElement(g, qn('a:srgbClr'))
    c.set('val', color)
    a = etree.SubElement(c, qn('a:alpha'))
    a.set('val', str(int(alpha * 1000)))


def soft_edge(shape, rad=0.1):
    eff = _effect_lst(_spPr(shape))
    se = etree.SubElement(eff, qn('a:softEdge'))
    se.set('rad', str(int(rad * 914400)))


def _ensure_ln_before_effects(shape):
    spPr = _spPr(shape)
    eff = spPr.find(qn('a:effectLst'))
    ln = spPr.find(qn('a:ln'))
    if eff is not None and ln is not None and spPr.index(ln) > spPr.index(eff):
        eff.addprevious(ln)


def grad_fill(shape, stops, angle=90, radial=False):
    """stops: list of (pos 0-100, color, alpha 0-100)."""
    spPr = _spPr(shape)
    for tag in ('a:solidFill', 'a:noFill', 'a:gradFill'):
        el = spPr.find(qn(tag))
        if el is not None:
            spPr.remove(el)
    g = etree.Element(qn('a:gradFill'))
    g.set('rotWithShape', '1')
    gl = etree.SubElement(g, qn('a:gsLst'))
    for pos, col, al in stops:
        gs = etree.SubElement(gl, qn('a:gs'))
        gs.set('pos', str(int(pos * 1000)))
        c = etree.SubElement(gs, qn('a:srgbClr'))
        c.set('val', col)
        if al < 100:
            a = etree.SubElement(c, qn('a:alpha'))
            a.set('val', str(int(al * 1000)))
    if radial:
        p = etree.SubElement(g, qn('a:path'))
        p.set('path', 'circle')
        f = etree.SubElement(p, qn('a:fillToRect'))
        for k in 'ltrb':
            f.set(k, '50000')
    else:
        lin = etree.SubElement(g, qn('a:lin'))
        lin.set('ang', str(int(angle * 60000)))
        lin.set('scaled', '0')
    geom = spPr.find(qn('a:prstGeom'))
    if geom is None:
        geom = spPr.find(qn('a:custGeom'))
    geom.addnext(g)


def name(shape, n):
    shape.name = n
    return shape


# ================================================================ primitives
def rect(s, x, y, w, h, fill=None, radius=None, line=None, lw=0.75, alpha=None, line_alpha=None,
         nm=None, shp=None, rot=0):
    kind = shp or (MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE)
    r = s.shapes.add_shape(kind, E(x), E(y), E(w), E(h))
    if radius and kind == MSO_SHAPE.ROUNDED_RECTANGLE:
        r.adjustments[0] = min(0.5, radius / min(w, h))
    set_fill(r, fill, alpha)
    set_line(r, line, lw, line_alpha)
    r.shadow.inherit = False
    if rot:
        r.rotation = rot
    tf = r.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    if nm:
        r.name = nm
    return r


def oval(s, x, y, w, h=None, fill=None, line=None, lw=0.75, alpha=None, line_alpha=None, nm=None, dash=None):
    h = w if h is None else h
    o = s.shapes.add_shape(MSO_SHAPE.OVAL, E(x), E(y), E(w), E(h))
    set_fill(o, fill, alpha)
    set_line(o, line, lw, line_alpha, dash)
    o.shadow.inherit = False
    if nm:
        o.name = nm
    return o


def line(s, x1, y1, x2, y2, color=None, width=1.0, alpha=None, dash=None, head=None, tail=None, nm=None):
    c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, E(x1), E(y1), E(x2), E(y2))
    c.line.color.rgb = rgb(color or T['line'])
    c.line.width = Pt(width)
    ln = c._element.spPr.find(qn('a:ln'))
    if alpha is not None:
        clr = ln.find(qn('a:solidFill'))[0]
        a = etree.SubElement(clr, qn('a:alpha'))
        a.set('val', str(int(alpha * 1000)))
    if dash:
        d = etree.SubElement(ln, qn('a:prstDash'))
        d.set('val', dash)
    for tag, kind in (('a:headEnd', head), ('a:tailEnd', tail)):
        if kind:
            e = etree.SubElement(ln, qn(tag))
            e.set('type', kind)
            e.set('w', 'med')
            e.set('len', 'med')
    etree.SubElement(c._element.spPr, qn('a:effectLst'))
    if nm:
        c.name = nm
    return c


def elbow(s, pts, color=None, width=1.0, dash=None, tail=None):
    """Polyline through points [(x,y),...] built from straight connectors."""
    out = []
    for i in range(len(pts) - 1):
        (x1, y1), (x2, y2) = pts[i], pts[i + 1]
        out.append(line(s, x1, y1, x2, y2, color, width, dash=dash,
                        tail=tail if i == len(pts) - 2 else None))
    return out


def freeform(s, pts, fill=None, line=None, lw=0.75, alpha=None, nm=None):
    """Closed polygon from absolute inch coordinates."""
    xs = [p[0] for p in pts]
    ys = [p[1] for p in pts]
    x0, y0 = min(xs), min(ys)
    scale = 914400
    fb = s.shapes.build_freeform(E(pts[0][0]), E(pts[0][1]), scale=1.0)
    fb.add_line_segments([(E(x), E(y)) for x, y in pts[1:]], close=True)
    shp = fb.convert_to_shape()
    set_fill(shp, fill, alpha)
    set_line(shp, line, lw)
    shp.shadow.inherit = False
    if nm:
        shp.name = nm
    return shp


# ================================================================ text
def _apply_run(run, text, font=F_BODY, size=14, color=None, bold=False, italic=False, spacing=None,
               outline=None, outline_w=1.0, caps=False, alpha=None):
    run.text = text
    f = run.font
    f.size = Pt(size)
    f.bold = bold
    f.italic = italic
    rPr = run._r.get_or_add_rPr()
    if outline:
        # hollow text: outline + no fill
        ln = etree.SubElement(rPr, qn('a:ln'))
        ln.set('w', str(int(outline_w * 12700)))
        sf = etree.SubElement(ln, qn('a:solidFill'))
        c = etree.SubElement(sf, qn('a:srgbClr'))
        c.set('val', outline)
        if alpha is not None:
            a = etree.SubElement(c, qn('a:alpha'))
            a.set('val', str(int(alpha * 1000)))
        etree.SubElement(rPr, qn('a:noFill'))
    elif color:
        sf = etree.SubElement(rPr, qn('a:solidFill'))
        c = etree.SubElement(sf, qn('a:srgbClr'))
        c.set('val', color)
        if alpha is not None:
            a = etree.SubElement(c, qn('a:alpha'))
            a.set('val', str(int(alpha * 1000)))
    for tag in ('a:latin', 'a:ea', 'a:cs'):
        el = etree.SubElement(rPr, qn(tag))
        el.set('typeface', font)
    if spacing is not None:
        rPr.set('spc', str(int(spacing * 100)))
    if caps:
        rPr.set('cap', 'all')


def text(s, x, y, w, h, content, size=14, color=None, font=F_BODY, bold=False, italic=False,
         align='l', anchor='t', spacing=None, line_sp=None, space_after=None, nm=None,
         outline=None, outline_w=1.0, vert=None, wrap=True, alpha=None, bullets=False,
         bullet_color=None, margin=0.0, autofit=False, on=None):
    """content: str | list of paragraphs; paragraph = str | list of runs; run = str | (str, {opts}).
    '\n' inside a str starts a new paragraph."""
    tb = s.shapes.add_textbox(E(x), E(y), E(w), E(h))
    tf = tb.text_frame
    tf.word_wrap = wrap
    from pptx.enum.text import MSO_AUTO_SIZE
    tf.auto_size = MSO_AUTO_SIZE.NONE
    if alpha is not None and not outline:
        base = on or (T['bg_dark'] if getattr(s, '_dark', False) else T['bg_light'])
        color = mix(color or T['text'], base, alpha / 100)
        alpha = None
    tf.margin_left = tf.margin_right = E(margin)
    tf.margin_top = tf.margin_bottom = E(margin)
    tf.vertical_anchor = {'t': MSO_ANCHOR.TOP, 'm': MSO_ANCHOR.MIDDLE, 'b': MSO_ANCHOR.BOTTOM}[anchor]
    bodyPr = tf._txBody.find(qn('a:bodyPr'))
    if vert:
        bodyPr.set('vert', vert)
    if autofit:
        etree.SubElement(bodyPr, qn('a:normAutofit'))
    if isinstance(content, str):
        paras = content.split('\n')
    else:
        paras = content
    color = color or T['text']
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = {'l': PP_ALIGN.LEFT, 'c': PP_ALIGN.CENTER, 'r': PP_ALIGN.RIGHT, 'j': PP_ALIGN.JUSTIFY}[align]
        if line_sp:
            p.line_spacing = line_sp
        if space_after is not None:
            p.space_after = Pt(space_after)
        runs = [para] if isinstance(para, (str, tuple)) else para
        for r in runs:
            if isinstance(r, str):
                t, o = r, {}
            else:
                t, o = r
            run = p.add_run()
            _apply_run(run, t, font=o.get('font', font), size=o.get('size', size), color=o.get('color', color),
                       bold=o.get('bold', bold), italic=o.get('italic', italic), spacing=o.get('spacing', spacing),
                       outline=o.get('outline', outline), outline_w=o.get('outline_w', outline_w),
                       alpha=alpha if o.get('outline', outline) else None)
        if bullets:
            pPr = p._p.get_or_add_pPr()
            pPr.set('marL', str(int(0.22 * 914400)))
            pPr.set('indent', str(int(-0.22 * 914400)))
            bc = etree.SubElement(pPr, qn('a:buClr'))
            c = etree.SubElement(bc, qn('a:srgbClr'))
            c.set('val', bullet_color or T['ink'])
            bs = etree.SubElement(pPr, qn('a:buSzPct'))
            bs.set('val', '100000')
            bf = etree.SubElement(pPr, qn('a:buFont'))
            bf.set('typeface', 'Arial')
            bu = etree.SubElement(pPr, qn('a:buChar'))
            bu.set('char', '•')
    if nm:
        tb.name = nm
    return tb


def shape_text(shp, content, size=12, color=None, font=F_SEMI, bold=False, align='c', anchor='m',
               spacing=None, margin=(0.1, 0.05), line_sp=None):
    tf = shp.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = E(margin[0])
    tf.margin_top = tf.margin_bottom = E(margin[1])
    tf.vertical_anchor = {'t': MSO_ANCHOR.TOP, 'm': MSO_ANCHOR.MIDDLE, 'b': MSO_ANCHOR.BOTTOM}[anchor]
    paras = content.split('\n') if isinstance(content, str) else content
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = {'l': PP_ALIGN.LEFT, 'c': PP_ALIGN.CENTER, 'r': PP_ALIGN.RIGHT}[align]
        if line_sp:
            p.line_spacing = line_sp
        runs = [para] if isinstance(para, (str, tuple)) else para
        for r in runs:
            t, o = (r, {}) if isinstance(r, str) else r
            _apply_run(p.add_run(), t, font=o.get('font', font), size=o.get('size', size),
                       color=o.get('color', color or T['ink']), bold=o.get('bold', bold), spacing=o.get('spacing', spacing))
    return shp


# ================================================================ pictures
CREDITS = []


def picture(s, path, x, y, w, h, radius=None, geom=None, focus=(0.5, 0.5), border=None, border_w=1.0,
            border_alpha=None, alpha=None, nm=None, fit='cover', rot=0):
    """Insert a picture that covers the box (cropped around focus), optional rounded/ellipse geometry."""
    from PIL import Image
    iw, ih = Image.open(path).size
    if fit == 'cover':
        pic = s.shapes.add_picture(path, E(x), E(y), E(w), E(h))
        box_r, img_r = w / h, iw / ih
        if img_r > box_r:          # too wide -> crop left/right
            keep = box_r / img_r
            extra = 1 - keep
            left = min(max(focus[0] - keep / 2, 0), extra)
            pic.crop_left, pic.crop_right = left, extra - left
        else:
            keep = img_r / box_r
            extra = 1 - keep
            top = min(max(focus[1] - keep / 2, 0), extra)
            pic.crop_top, pic.crop_bottom = top, extra - top
    else:                          # contain: keep aspect inside the box, centred at bottom
        r = min(w / iw, h / ih)
        pw, ph = iw * r, ih * r
        pic = s.shapes.add_picture(path, E(x + (w - pw) / 2), E(y + h - ph), E(pw), E(ph))
    spPr = pic._element.spPr
    if radius or geom:
        g = spPr.find(qn('a:prstGeom'))
        g.set('prst', geom or 'roundRect')
        av = g.find(qn('a:avLst'))
        if av is None:
            av = etree.SubElement(g, qn('a:avLst'))
        if radius and (geom in (None, 'roundRect')):
            gd = etree.SubElement(av, qn('a:gd'))
            gd.set('name', 'adj')
            gd.set('fmla', f'val {int(min(50000, radius / min(w, h) * 100000))}')
    if border:
        set_line(pic, border, border_w, border_alpha)
    if alpha is not None:
        blip = pic._element.find('.//' + qn('a:blip'))
        am = etree.SubElement(blip, qn('a:alphaModFix'))
        am.set('amt', str(int(alpha * 1000)))
    if rot:
        pic.rotation = rot
    if nm:
        pic.name = nm
    return pic


def credit(img_id, where):
    CREDITS.append((img_id, where))


# ================================================================ icons
ICON_DIR = os.path.join(ROOT, 'assets', 'icons', 'png')
LUCIDE = os.path.join(ROOT, 'assets', 'icons', 'svg')   # Lucide (ISC licence) SVG set


def icon_png(nm, color, px=256, stroke=2.0):
    os.makedirs(ICON_DIR, exist_ok=True)
    out = os.path.join(ICON_DIR, f'{nm}_{color}.png')
    if os.path.exists(out):
        return out
    import cairosvg
    src = os.path.join(LUCIDE, f'{nm}.svg')
    svg = open(src, encoding='utf-8').read()
    svg = svg.replace('currentColor', '#' + color).replace('stroke-width="2"', f'stroke-width="{stroke}"')
    cairosvg.svg2png(bytestring=svg.encode(), write_to=out, output_width=px, output_height=px)
    return out


def icon(s, nm, x, y, size, color, stroke=2.0, nmx=None):
    p = s.shapes.add_picture(icon_png(nm, color, stroke=stroke), E(x), E(y), E(size), E(size))
    if nmx:
        p.name = nmx
    return p


def icon_badge(s, nm, cx, cy, d, fill, color, ring=None, ring_gap=0.08, icon_scale=0.5, shadow_on=False):
    """Icon centred in a filled circle, optional outer ring."""
    shapes = []
    if ring:
        shapes.append(oval(s, cx - d / 2 - ring_gap, cy - d / 2 - ring_gap, d + 2 * ring_gap, fill=None,
                           line=ring, lw=1.0, dash='dash'))
    c = oval(s, cx - d / 2, cy - d / 2, d, fill=fill)
    if shadow_on:
        shadow(c, blur=0.2, dist=0.05, alpha=20)
    shapes.append(c)
    isz = d * icon_scale
    shapes.append(icon(s, nm, cx - isz / 2, cy - isz / 2, isz, color))
    return shapes


# ================================================================ decorative background
BLOB_POS = [
    # (x, y, d) centres in inches for 3 blobs; picked per slide to create drift under Morph
    [(12.6, -0.6, 4.6), (-0.8, 7.6, 3.8), (13.4, 6.4, 3.0)],
    [(11.8, 0.2, 4.2), (0.2, 7.9, 4.2), (13.6, 4.8, 2.8)],
    [(13.2, 1.0, 5.0), (-0.4, 6.8, 3.4), (7.2, 8.4, 3.2)],
    [(12.2, -0.9, 4.0), (0.6, 8.2, 4.6), (14.0, 7.2, 3.6)],
    [(13.8, 0.4, 4.4), (-1.0, 5.2, 3.6), (9.6, 8.6, 3.0)],
    [(11.4, -1.0, 4.8), (1.2, 8.4, 3.8), (13.9, 5.6, 3.2)],
]


GLOW_DIR = os.path.join(ROOT, 'assets', 'glow')


def glow_png(color, alpha, px=512, falloff=1.7):
    """Soft radial glow as RGBA PNG (cached)."""
    os.makedirs(GLOW_DIR, exist_ok=True)
    out = os.path.join(GLOW_DIR, f'glow_{color}_{int(alpha)}_{falloff}.png')
    if not os.path.exists(out):
        import numpy as np
        from PIL import Image
        yy, xx = np.mgrid[0:px, 0:px]
        r = np.sqrt((xx - px / 2 + 0.5) ** 2 + (yy - px / 2 + 0.5) ** 2) / (px / 2)
        r = np.clip(r, 0, 1)
        a = (1 - (3 * r ** 2 - 2 * r ** 3)) ** falloff * (alpha / 100)
        rgbv = [int(color[i:i + 2], 16) for i in (0, 2, 4)]
        arr = np.zeros((px, px, 4), dtype=np.uint8)
        arr[..., 0], arr[..., 1], arr[..., 2] = rgbv
        arr[..., 3] = (a * 255).astype(np.uint8)
        Image.fromarray(arr, 'RGBA').save(out)
    return out


def glow_img(s, cx, cy, d, color, alpha, nm=None, falloff=1.7):
    p = s.shapes.add_picture(glow_png(color, alpha, falloff=falloff), E(cx - d / 2), E(cy - d / 2), E(d), E(d))
    if nm:
        p.name = nm
    return p


def blobs(s, idx=0, dark=None, strength=1.0, pos=None):
    dark = s._dark if dark is None else dark
    cols = T['blob_dark'] if dark else T['blob_light']
    pos = pos or BLOB_POS[idx % len(BLOB_POS)]
    out = []
    for i, ((cx, cy, d), (col, al)) in enumerate(zip(pos, cols)):
        out.append(glow_img(s, cx, cy, d * 1.35, col, al * strength, nm=f'!!blob{i + 1}'))
    return out


def page_no(s, n, dark=None):
    dark = s._dark if dark is None else dark
    text(s, SW - ML - 0.8, SH - 0.42, 0.8, 0.25, f'{n:02d}', size=9, font=F_SEMI,
         color='FFFFFF' if dark else T['muted'], align='r', alpha=55 if dark else None)


def dot_line(s, x, y, w, color=None, dot=0.09, width=1.0, alpha=None):
    color = color or T['ink']
    line(s, x, y, x + w, y, color, width, alpha=alpha)
    oval(s, x + w - dot / 2, y - dot / 2, dot, fill=color)


# ================================================================ components
def breadcrumb(s, num, label, dark=None):
    dark = s._dark if dark is None else dark
    col = 'FFFFFF' if dark else T['ink']
    a = text(s, ML, 0.34, 0.45, 0.28, f'{num:02d}', size=10, font=F_HEAD, color=T['accent'] if dark else T['ink'],
             nm=f'!!num{num:02d}')
    b = text(s, ML + 0.42, 0.34, 4.5, 0.28, label.upper(), size=10, font=F_SEMI, color=col,
             spacing=1.0, alpha=70 if dark else None, nm=f'!!sec{num:02d}')
    return a, b


def title(s, x, y, w, h, t, size=34, color=None, align='l', anchor='t', nm=None, line_sp=0.95, max_lines=None):
    if max_lines is None:
        max_lines = max(t.count('\n') + 1, int(h / (size / 72 * 1.05)))
    size = fit_size(t, F_HEAD, size, w, max_lines)
    lines = n_lines(t, F_HEAD, size, w * 0.97)
    h = max(h, lines * size / 72 * 1.05)
    return text(s, x, y, w, h, t, size=size, font=F_HEAD, color=color or ('FFFFFF' if s._dark else T['ink']),
                align=align, anchor=anchor, nm=nm, line_sp=line_sp)


def subhead(s, x, y, w, t, size=15, color=None, bar=True):
    """Section sub-heading with a short vertical accent bar on its left (reference style)."""
    col = color or ('FFFFFF' if s._dark else T['ink'])
    out = []
    if bar:
        out.append(rect(s, x, y + 0.04, 0.05, 0.3, fill=T['accent'] if s._dark else T['ink']))
    out.append(text(s, x + (0.16 if bar else 0), y, w, 0.38, t.upper(), size=size, font=F_HEAD, color=col, anchor='m'))
    return out


def pill(s, x, y, t, size=11, fill=None, color=None, h=0.34, w=None, font=F_SEMI, pad=0.22, nm=None):
    fill = fill or T['accent']
    color = color or T['accent_ink']
    if w is None:
        w = max(0.9, len(t) * size * 0.0092 + 2 * pad)
    r = rect(s, x, y, w, h, fill=fill, radius=h / 2, nm=nm)
    shape_text(r, t, size=size, color=color, font=font, margin=(0.08, 0.0))
    return r


def card(s, x, y, w, h, fill=None, radius=0.22, shadow_on=True, line=None, alpha=None, nm=None, sh_alpha=16):
    r = rect(s, x, y, w, h, fill=fill or T['card'], radius=radius, line=line, alpha=alpha, nm=nm)
    if shadow_on:
        shadow(r, blur=0.35, dist=0.12, alpha=sh_alpha, color=T['ink'])
    return r


def pin(s, x, y, size=0.42):
    """Push-pin accent (reference motif) — icon on accent disc."""
    icon_badge(s, 'pin', x + size / 2, y + size / 2, size, T['accent'], T['accent_ink'], icon_scale=0.58)


def stat(s, x, y, value, unit='', size=40, color=None, unit_size=None, font=F_HEAD, w=3.0, h=0.8, align='l'):
    col = color or T['ink']
    runs = [(value, {'size': size, 'color': col})]
    if unit:
        runs.append((' ' + unit, {'size': unit_size or size * 0.42, 'color': col, 'font': F_HEAD}))
    return text(s, x, y, w, h, [runs], size=size, font=font, color=col, align=align, anchor='b')


# ================================================================ charts
def _chart_text(chart, size=10, color=None):
    chart.has_title = False
    chart.font.size = Pt(size)
    chart.font.name = F_BODY
    chart.font.color.rgb = rgb(color or T['muted'])


def _no_fill_chart(chart):
    cs = chart._chartSpace
    spPr = cs.find(qn('c:spPr'))
    if spPr is None:
        spPr = etree.SubElement(cs, qn('c:spPr'))
        # c:spPr must precede c:txPr / externalData etc.
        txPr = cs.find(qn('c:txPr'))
        if txPr is not None:
            txPr.addprevious(spPr)
    for ch in list(spPr):
        spPr.remove(ch)
    etree.SubElement(spPr, qn('a:noFill'))
    ln = etree.SubElement(spPr, qn('a:ln'))
    etree.SubElement(ln, qn('a:noFill'))


def _axis_clean(ax, show_line=False, color=None, size=10, grid=False, grid_color=None):
    ax.has_major_gridlines = grid
    if grid:
        gl = ax.major_gridlines.format.line
        gl.color.rgb = rgb(grid_color or T['line'])
        gl.width = Pt(0.6)
        gl.dash_style = None
    ax.format.line.fill.background() if not show_line else None
    if show_line:
        ax.format.line.color.rgb = rgb(T['line'])
        ax.format.line.width = Pt(0.75)
    ax.tick_labels.font.size = Pt(size)
    ax.tick_labels.font.color.rgb = rgb(color or T['muted'])
    ax.tick_labels.font.name = F_BODY


def chart_bar(s, x, y, w, h, cats, series, horizontal=False, colors=None, gap=80, overlap=-10,
              labels=False, num_fmt='0', legend=True, dark=False, label_color=None, max_v=None, font_size=10,
              legend_pos='b'):
    cd = CategoryChartData()
    cd.categories = cats
    for nm_, vals in series:
        cd.add_series(nm_, vals)
    kind = XL_CHART_TYPE.BAR_CLUSTERED if horizontal else XL_CHART_TYPE.COLUMN_CLUSTERED
    gf = s.shapes.add_chart(kind, E(x), E(y), E(w), E(h), cd)
    ch = gf.chart
    tcol = 'C9D3EA' if dark else T['muted']
    _chart_text(ch, font_size, tcol)
    _no_fill_chart(ch)
    colors = colors or T['chart']
    plot = ch.plots[0]
    plot.gap_width = gap
    plot.overlap = overlap
    for i, ser in enumerate(plot.series):
        ser.format.fill.solid()
        ser.format.fill.fore_color.rgb = rgb(colors[i % len(colors)])
        ser.invert_if_negative = False
    _axis_clean(ch.category_axis, show_line=True, color=tcol, size=font_size)
    _axis_clean(ch.value_axis, color=tcol, size=font_size, grid=True, grid_color=T['grid_dark'] if dark else T['line'])
    ch.value_axis.tick_labels.number_format = num_fmt
    ch.value_axis.tick_labels.number_format_is_linked = False
    if max_v:
        ch.value_axis.maximum_scale = max_v
    ch.value_axis.minimum_scale = 0
    if labels:
        plot.has_data_labels = True
        dl = plot.data_labels
        dl.font.size = Pt(font_size)
        dl.font.bold = True
        dl.font.color.rgb = rgb(label_color or (T['accent'] if dark else T['ink']))
        dl.number_format = num_fmt
        dl.number_format_is_linked = False
        dl.position = XL_LABEL_POSITION.OUTSIDE_END
    ch.has_legend = legend and len(series) > 1
    if ch.has_legend:
        ch.legend.position = {'b': XL_LEGEND_POSITION.BOTTOM, 'r': XL_LEGEND_POSITION.RIGHT, 't': XL_LEGEND_POSITION.TOP}[legend_pos]
        ch.legend.include_in_layout = False
        ch.legend.font.size = Pt(font_size)
        ch.legend.font.color.rgb = rgb(tcol)
    return gf


def chart_line(s, x, y, w, h, cats, series, colors=None, smooth=True, width=2.5, legend=True, dark=False,
               num_fmt='0', markers=False, font_size=10, area=False, max_v=None, legend_pos='r'):
    cd = CategoryChartData()
    cd.categories = cats
    for nm_, vals in series:
        cd.add_series(nm_, vals)
    kind = XL_CHART_TYPE.AREA if area else (XL_CHART_TYPE.LINE_MARKERS if markers else XL_CHART_TYPE.LINE)
    gf = s.shapes.add_chart(kind, E(x), E(y), E(w), E(h), cd)
    ch = gf.chart
    tcol = 'C9D3EA' if dark else T['muted']
    _chart_text(ch, font_size, tcol)
    _no_fill_chart(ch)
    colors = colors or T['chart']
    for i, ser in enumerate(ch.plots[0].series):
        col = colors[i % len(colors)]
        if area:
            spPr = ser._element.get_or_add_spPr()
            for ch_ in list(spPr):
                spPr.remove(ch_)
            g = X(f'<a:gradFill rotWithShape="1"><a:gsLst><a:gs pos="0"><a:srgbClr val="{col}"><a:alpha val="95000"/></a:srgbClr></a:gs>'
                  f'<a:gs pos="100000"><a:srgbClr val="{col}"><a:alpha val="15000"/></a:srgbClr></a:gs></a:gsLst><a:lin ang="5400000" scaled="0"/></a:gradFill>')
            spPr.append(g)
            ln = X(f'<a:ln w="{int(width * 12700)}"><a:solidFill><a:srgbClr val="{col}"/></a:solidFill></a:ln>')
            spPr.append(ln)
        else:
            ser.format.line.color.rgb = rgb(col)
            ser.format.line.width = Pt(width)
            ser.smooth = smooth
            if markers:
                ser.marker.style = XL_MARKER_STYLE.CIRCLE
                ser.marker.size = 7
                ser.marker.format.fill.solid()
                ser.marker.format.fill.fore_color.rgb = rgb(col)
                ser.marker.format.line.color.rgb = rgb('FFFFFF')
            else:
                ser.marker.style = XL_MARKER_STYLE.NONE
    _axis_clean(ch.category_axis, show_line=True, color=tcol, size=font_size)
    _axis_clean(ch.value_axis, color=tcol, size=font_size, grid=True, grid_color=T['grid_dark'] if dark else T['line'])
    ch.value_axis.tick_labels.number_format = num_fmt
    ch.value_axis.tick_labels.number_format_is_linked = False
    if max_v:
        ch.value_axis.maximum_scale = max_v
    ch.value_axis.minimum_scale = 0
    ch.has_legend = legend and len(series) > 1
    if ch.has_legend:
        ch.legend.position = {'b': XL_LEGEND_POSITION.BOTTOM, 'r': XL_LEGEND_POSITION.RIGHT, 't': XL_LEGEND_POSITION.TOP}[legend_pos]
        ch.legend.include_in_layout = False
        ch.legend.font.size = Pt(font_size)
        ch.legend.font.color.rgb = rgb(tcol)
    return gf


def chart_doughnut(s, x, y, w, h, cats, vals, colors=None, hole=68, labels=True, font_size=10, legend=False):
    cd = CategoryChartData()
    cd.categories = cats
    cd.add_series('S', vals)
    gf = s.shapes.add_chart(XL_CHART_TYPE.DOUGHNUT, E(x), E(y), E(w), E(h), cd)
    ch = gf.chart
    _chart_text(ch, font_size, T['muted'])
    _no_fill_chart(ch)
    colors = colors or T['chart']
    plot = ch.plots[0]
    ser = plot.series[0]
    for i, pt in enumerate(ser.points):
        pt.format.fill.solid()
        pt.format.fill.fore_color.rgb = rgb(colors[i % len(colors)])
        pt.format.line.color.rgb = rgb(T['bg_light'])
        pt.format.line.width = Pt(2.5)
    dn = ch._chartSpace.find('.//' + qn('c:doughnutChart'))
    hs = dn.find(qn('c:holeSize'))
    if hs is None:
        hs = etree.SubElement(dn, qn('c:holeSize'))
    hs.set('val', str(hole))
    if labels:
        plot.has_data_labels = True
        dl = plot.data_labels
        dl.number_format = '0%'
        dl.number_format_is_linked = False
        dl.show_percentage = True
        dl.show_value = False
        dl.font.size = Pt(font_size)
        dl.font.bold = True
        dl.font.color.rgb = rgb('FFFFFF')
    ch.has_legend = legend
    return gf


# ================================================================ tables
def table(s, x, y, w, h, rows, col_w=None, header_fill=None, header_color=None, body_fill=None, body_fill2=None,
          body_color=None, line_color=None, size=11, header_size=11, first_col_bold=True, align_num='c', row_h=None):
    nr, nc = len(rows), len(rows[0])
    gf = s.shapes.add_table(nr, nc, E(x), E(y), E(w), E(h))
    tbl = gf.table
    # drop the default table style banding
    tblPr = tbl._tbl.tblPr
    tblPr.set('firstRow', '0')
    tblPr.set('bandRow', '0')
    sid = tblPr.find(qn('a:tableStyleId'))
    if sid is not None:
        sid.text = '{2D5ABB26-0587-4C30-8999-92F81FD0307C}'   # "No Style, No Grid"
    if col_w:
        for i, cw in enumerate(col_w):
            tbl.columns[i].width = E(cw)
    rh = row_h or h / nr
    for r in range(nr):
        tbl.rows[r].height = E(rh)
    for r in range(nr):
        for c in range(nc):
            cell = tbl.cell(r, c)
            hdr = r == 0
            fill = header_fill if hdr else (body_fill2 if (body_fill2 and r % 2 == 0) else body_fill)
            if fill:
                cell.fill.solid()
                cell.fill.fore_color.rgb = rgb(fill)
            else:
                cell.fill.background()
            cell.margin_left = cell.margin_right = E(0.1)
            cell.margin_top = cell.margin_bottom = E(0.03)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.LEFT if (c == 0 and not hdr) else (PP_ALIGN.CENTER if align_num == 'c' else PP_ALIGN.RIGHT)
            run = p.add_run()
            is_bold = hdr or (c == 0 and first_col_bold)
            _apply_run(run, str(rows[r][c]), font=F_SEMI if is_bold else F_BODY,
                       size=header_size if hdr else size,
                       color=(header_color if hdr else body_color) or T['ink'])
            # borders
            tcPr = cell._tc.get_or_add_tcPr()
            for tag in ('a:lnL', 'a:lnR', 'a:lnT', 'a:lnB'):
                ln = etree.SubElement(tcPr, qn(tag))
                ln.set('w', str(int(0.75 * 12700)))
                sf = etree.SubElement(ln, qn('a:solidFill'))
                cc = etree.SubElement(sf, qn('a:srgbClr'))
                cc.set('val', line_color or T['line'])
            # fill must come after borders in tcPr
            for ftag in ('a:solidFill', 'a:noFill'):
                f = tcPr.find(qn(ftag))
                if f is not None:
                    tcPr.remove(f)
                    tcPr.append(f)
    return gf


# ================================================================ animation + transitions
EFFECTS = {
    # name: (presetID, subtype)
    'fade': (10, 0), 'wipe_l': (22, 8), 'wipe_u': (22, 4), 'wipe_d': (22, 1), 'wipe_r': (22, 2),
    'zoom': (53, 16), 'float': (42, 0), 'rise': (42, 0), 'wheel': (21, 1), 'fly_l': (2, 8), 'fly_r': (2, 2),
    'fly_b': (2, 4), 'grow': (53, 16),
}


def anim(s, shape, effect='fade', delay=0.0, dur=0.5):
    """Queue an entrance animation (auto-plays after the slide transition)."""
    s._anims.append((shape, effect, int(delay * 1000), int(dur * 1000)))
    return shape


def anims(s, shapes, effect='fade', start=0.0, step=0.12, dur=0.5):
    t = start
    for sh in shapes:
        if isinstance(sh, (list, tuple)):
            for x in sh:
                anim(s, x, effect, t, dur)
        else:
            anim(s, sh, effect, t, dur)
        t += step
    return t


def morph(s, dur=1.2):
    s._morph = ('morph', dur)


def fade_tr(s, dur=0.7):
    s._morph = ('fade', dur)


def _tgt(spid):
    return f'<p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl>'


def _bhv(ids, spid, effect, dur):
    """Return XML for the behaviour children of one effect."""
    x = []

    def nid():
        ids[0] += 1
        return ids[0]

    x.append(f'<p:set><p:cBhvr><p:cTn id="{nid()}" dur="1" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst></p:cTn>'
             f'{_tgt(spid)}<p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr>'
             f'<p:to><p:strVal val="visible"/></p:to></p:set>')

    def anim_prop(prop, v0, v1):
        return (f'<p:anim calcmode="lin" valueType="num"><p:cBhvr additive="base"><p:cTn id="{nid()}" dur="{dur}" fill="hold"/>'
                f'{_tgt(spid)}<p:attrNameLst><p:attrName>{prop}</p:attrName></p:attrNameLst></p:cBhvr>'
                f'<p:tavLst><p:tav tm="0"><p:val><p:strVal val="{v0}"/></p:val></p:tav>'
                f'<p:tav tm="100000"><p:val><p:strVal val="{v1}"/></p:val></p:tav></p:tavLst></p:anim>')

    def fx(filt):
        return (f'<p:animEffect transition="in" filter="{filt}"><p:cBhvr><p:cTn id="{nid()}" dur="{dur}"/>'
                f'{_tgt(spid)}</p:cBhvr></p:animEffect>')

    if effect == 'fade':
        x.append(fx('fade'))
    elif effect.startswith('wipe'):
        d = {'wipe_l': 'left', 'wipe_u': 'down', 'wipe_d': 'up', 'wipe_r': 'right'}[effect]
        x.append(fx(f'wipe({d})'))
    elif effect in ('zoom', 'grow'):
        k = '0.7' if effect == 'zoom' else '0.3'
        x.append(anim_prop('ppt_w', f'#ppt_w*{k}', '#ppt_w'))
        x.append(anim_prop('ppt_h', f'#ppt_h*{k}', '#ppt_h'))
        x.append(fx('fade'))
    elif effect in ('float', 'rise'):
        x.append(fx('fade'))
        x.append(anim_prop('ppt_x', '#ppt_x', '#ppt_x'))
        x.append(anim_prop('ppt_y', '#ppt_y+.06', '#ppt_y'))
    elif effect == 'wheel':
        x.append(fx('wheel(1)'))
    elif effect.startswith('fly'):
        if effect == 'fly_l':
            x.append(anim_prop('ppt_x', '0-#ppt_w/2', '#ppt_x'))
            x.append(anim_prop('ppt_y', '#ppt_y', '#ppt_y'))
        elif effect == 'fly_r':
            x.append(anim_prop('ppt_x', '1+#ppt_w/2', '#ppt_x'))
            x.append(anim_prop('ppt_y', '#ppt_y', '#ppt_y'))
        else:
            x.append(anim_prop('ppt_x', '#ppt_x', '#ppt_x'))
            x.append(anim_prop('ppt_y', '1+#ppt_h/2', '#ppt_y'))
    return ''.join(x)


def finalize(s):
    """Write transition + timing XML for a slide."""
    sld = s._element
    P = NS['p']
    for er in sld.iter(qn('a:effectRef')):
        er.set('idx', '0')
    # ---- transition
    if s._morph:
        kind, dur = s._morph
        ms = int(dur * 1000)
        if kind == 'morph':
            xml = (f'<mc:AlternateContent xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006">'
                   f'<mc:Choice xmlns:p159="http://schemas.microsoft.com/office/powerpoint/2015/09/main" Requires="p159">'
                   f'<p:transition xmlns:p="{P}" xmlns:p14="http://schemas.microsoft.com/office/powerpoint/2010/main" spd="slow" p14:dur="{ms}">'
                   f'<p159:morph option="byObject"/></p:transition></mc:Choice>'
                   f'<mc:Fallback><p:transition xmlns:p="{P}" spd="slow"><p:fade/></p:transition></mc:Fallback></mc:AlternateContent>')
        else:
            xml = (f'<mc:AlternateContent xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006">'
                   f'<mc:Choice xmlns:p14="http://schemas.microsoft.com/office/powerpoint/2010/main" Requires="p14">'
                   f'<p:transition xmlns:p="{P}" spd="slow" p14:dur="{ms}"><p:fade/></p:transition></mc:Choice>'
                   f'<mc:Fallback><p:transition xmlns:p="{P}" spd="slow"><p:fade/></p:transition></mc:Fallback></mc:AlternateContent>')
        tr = etree.fromstring(xml)
        clr = sld.find(qn('p:clrMapOvr'))
        clr.addnext(tr)
    # ---- timing
    if not s._anims:
        return
    ids = [4]
    effects_xml = []
    bld = []
    seen = set()
    for i, (shape, effect, delay, dur) in enumerate(sorted(s._anims, key=lambda a: a[2])):
        spid = shape.shape_id
        pid, sub = EFFECTS[effect]
        ids[0] += 1
        ctn = ids[0]
        node = 'afterEffect' if i == 0 else 'withEffect'
        body = _bhv(ids, spid, effect, dur)
        effects_xml.append(
            f'<p:par><p:cTn id="{ctn}" presetID="{pid}" presetClass="entr" presetSubtype="{sub}" fill="hold" grpId="0" nodeType="{node}">'
            f'<p:stCondLst><p:cond delay="{delay}"/></p:stCondLst><p:childTnLst>{body}</p:childTnLst></p:cTn></p:par>')
        tag = etree.QName(shape._element).localname
        if spid not in seen:
            seen.add(spid)
            if tag == 'sp':
                bld.append(f'<p:bldP spid="{spid}" grpId="0" animBg="1"/>')
            elif tag == 'graphicFrame':
                bld.append(f'<p:bldGraphic spid="{spid}" grpId="0"><p:bldAsOne/></p:bldGraphic>')
    timing = (f'<p:timing xmlns:p="{P}"><p:tnLst><p:par><p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot"><p:childTnLst>'
              f'<p:seq concurrent="1" nextAc="seek"><p:cTn id="2" dur="indefinite" nodeType="mainSeq"><p:childTnLst>'
              f'<p:par><p:cTn id="3" fill="hold"><p:stCondLst><p:cond delay="indefinite"/><p:cond evt="onBegin" delay="0"><p:tn val="2"/></p:cond></p:stCondLst><p:childTnLst>'
              f'<p:par><p:cTn id="4" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>'
              f'{"".join(effects_xml)}'
              f'</p:childTnLst></p:cTn></p:par>'
              f'</p:childTnLst></p:cTn></p:par>'
              f'</p:childTnLst></p:cTn><p:prevCondLst><p:cond evt="onPrev" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:prevCondLst>'
              f'<p:nextCondLst><p:cond evt="onNext" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:nextCondLst></p:seq>'
              f'</p:childTnLst></p:cTn></p:par></p:tnLst>'
              f'{"<p:bldLst>" + "".join(bld) + "</p:bldLst>" if bld else ""}</p:timing>')
    tm = etree.fromstring(timing)
    ext = sld.find(qn('p:extLst'))
    if ext is not None:
        ext.addprevious(tm)
    else:
        sld.append(tm)


def notes(s, t):
    s.notes_slide.notes_text_frame.text = t


# ================================================================ extra components
def s_curve(s, x1, y1, x2, y2, color=None, width=1.5, dash=None):
    """Curved connector with arrow head (cycle diagrams)."""
    c = s.shapes.add_connector(MSO_CONNECTOR.CURVE, E(x1), E(y1), E(x2), E(y2))
    c.line.color.rgb = rgb(color or T['ink'])
    c.line.width = Pt(width)
    ln = c._element.spPr.find(qn('a:ln'))
    if dash:
        d = etree.SubElement(ln, qn('a:prstDash'))
        d.set('val', dash)
    e = etree.SubElement(ln, qn('a:tailEnd'))
    e.set('type', 'triangle')
    e.set('w', 'med')
    e.set('len', 'med')
    return c


class _Box:
    """Adapter so helper functions can draw into a group shape."""
    def __init__(self, grp, dark=False):
        self.shapes = grp.shapes
        self._dark = dark


def polaroid(s, path, x, y, w, h, caption, rot=0):
    grp = s.shapes.add_group_shape()
    g = _Box(grp, s._dark)
    fr = rect(g, x, y, w, h, fill='FFFFFF')
    shadow(fr, blur=0.35, dist=0.12, alpha=45, color='000000', angle=100)
    m = 0.14
    picture(g, path, x + m, y + m, w - 2 * m, h - 0.72)
    text(g, x + m, y + h - 0.56, w - 2 * m, 0.42, caption, size=11, font=F_SEMI, italic=True, color=T['ink'],
         align='c', anchor='m')
    xfrm = grp._element.grpSpPr.find(qn('a:xfrm'))
    xfrm.set('rot', str(int(rot * 60000)))
    return grp


def glass(s, blurred_path, x, y, w, h, radius=0.3, tint=None, tint_alpha=30):
    """Frosted-glass card: the blurred background re-cropped to the card area + translucent tint and hairline."""
    pic = s.shapes.add_picture(blurred_path, E(x), E(y), E(w), E(h))
    pic.crop_left = x / SW
    pic.crop_right = 1 - (x + w) / SW
    pic.crop_top = y / SH
    pic.crop_bottom = 1 - (y + h) / SH
    g = pic._element.spPr.find(qn('a:prstGeom'))
    g.set('prst', 'roundRect')
    av = g.find(qn('a:avLst'))
    if av is None:
        av = etree.SubElement(g, qn('a:avLst'))
    gd = etree.SubElement(av, qn('a:gd'))
    gd.set('name', 'adj')
    gd.set('fmla', f'val {int(min(50000, radius / min(w, h) * 100000))}')
    ov = rect(s, x, y, w, h, fill=tint or T['bg_dark'], radius=radius, alpha=tint_alpha, line='FFFFFF', lw=1.25,
              line_alpha=45)
    shadow(ov, blur=0.5, dist=0.15, alpha=35, color='000000')
    return [pic, ov]
