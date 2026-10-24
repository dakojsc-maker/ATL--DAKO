"""
Product-showcase layouts shared by M06 "Thiên nhiên" (TH true MILK style) and M07 "Nông sản" (Highlands style).

Look & feel (both modelled on brand-website style decks)
    * a website navigation bar on every content slide: logo, menu with a moving underline, search / user / bag icons
    * cover: a huge brand word behind the product cut-outs, script word on top, flying decorations
    * product line-up cards, giant category words with the product in front of (M06) or inside (M07) the word
    * product carousel: one slide per product; the hero product sits in the middle, the others float small and
      blurred around it, name + script variant bottom-left, thumbnails bottom-right. Every product keeps the same
      shape name on every carousel slide, so Morph slides the hero away and brings the next one forward.
    * story cards with faded sepia photos (M06) or story slides with a big product and blurred decorations (M07)
    * features around the product, spec table, usage steps, results chart, browser-window cards, contact slide
Theme switches palette, fonts, title treatment and backgrounds: register('thiennhien' | 'nongsan').
Each layout: fn(prs, ctx, c) -> slide or [slides].
ctx = dict(n, img{photo key: path}, prod{product key: dict(path, blur)}, art{key: path}, brand{name, icon, menu})
"""
import re

import pptkit as K
from pptx.enum.shapes import MSO_SHAPE

SW, SH, ML = K.SW, K.SH, 0.6

THEMES = {
    'thiennhien': dict(
        name='Thiên nhiên',
        bg='FFFFFF', ink='3B2A20', text='5B4E45', muted='9A8F86', title='7A3B22', accent='3E8E3A', accent2='A8D26B',
        soft='F2F5EE', line='DDE5D6', card='FFFFFF', nav='FFFFFF', chart=['3E8E3A', 'A8D26B', '7A3B22', 'E2B34A'],
        fonts=dict(title=('Oswald', 'thiennhien/Oswald-Bold.ttf', True),
                   script=('Great Vibes', 'thiennhien/GreatVibes-Regular.ttf', False),
                   body=('Nunito Sans', 'thiennhien/NunitoSans-Regular.ttf', False),
                   semi=('Nunito Sans SemiBold', 'thiennhien/NunitoSans-SemiBold.ttf', False),
                   bold=('Nunito Sans', 'thiennhien/NunitoSans-Bold.ttf', True)),
        title_style='extrude', word_style='front', paper=False),
    'nongsan': dict(
        name='Nông sản',
        bg='F6EFE3', ink='3B2416', text='5C4636', muted='9C8A78', title='B5172F', accent='B5172F', accent2='E8A93A',
        soft='FBF6EE', line='E8DCC8', card='FFFDF8', nav='FBF6EE', chart=['B5172F', 'E8A93A', '6E8B3D', '3B2416'],
        fonts=dict(title=('Baloo 2 ExtraBold', 'nongsan/Baloo2-ExtraBold.ttf', False),
                   script=('Dancing Script', 'nongsan/DancingScript-Bold.ttf', True),
                   body=('Mulish', 'nongsan/Mulish-Regular.ttf', False),
                   semi=('Mulish SemiBold', 'nongsan/Mulish-SemiBold.ttf', False),
                   bold=('Mulish ExtraBold', 'nongsan/Mulish-ExtraBold.ttf', False)),
        title_style='flat', word_style='gap', paper=True),
}
T = THEMES['thiennhien']


def F(role):
    return T['fonts'][role][0]


def BOLD():
    """Bold flag needed for the 'bold' body role (a face whose bold is a style of the regular family)."""
    return T['fonts']['bold'][2]


def register(key):
    global T
    T = THEMES[key]
    K.THEMES['showcase_' + key] = dict(
        name=T['name'], key='showcase_' + key, bg_light=T['bg'], bg_dark='1E1A16', ink=T['ink'], text=T['text'],
        muted=T['muted'], soft=T['soft'], card=T['card'], card2=T['soft'], accent=T['accent'], accent_ink='FFFFFF',
        violet=T['accent2'], hot=T['title'], up=T['accent'], line=T['line'], chart=T['chart'],
        blob_light=[(T['accent2'], 25)], blob_dark=[(T['accent2'], 25)], tint='1E1A16', grid_dark='444444',
        row_alt=T['soft'])
    K.set_theme('showcase_' + key)
    for role in ('title', 'script', 'body', 'semi', 'bold'):      # first file wins for a face (regular before bold)
        face, fname, bold = T['fonts'][role]
        if role == 'bold' and face in K.FONT_FILES and K.FONT_FILES[face] != fname:
            continue
        K.FONT_FILES[face] = fname
        if bold and role in ('title', 'script'):
            K.BOLD_FACES.add(face)
    return T


# ------------------------------------------------------------------ helpers
def _new(prs, ctx, tr=1.2):
    ctx['n'] += 1
    s = K.blank_slide(prs, False)
    if ctx['n'] > 1:
        K.morph(s, tr)
    return s


def pic(ctx, key):
    if key in ctx['prod']:
        return ctx['prod'][key]['path']
    if key in ctx['img']:
        return ctx['img'][key]
    if key in ctx['art']:
        return ctx['art'][key]
    import content_io as CIO
    return CIO.photo_path(key)


def blur(ctx, key):
    return ctx['prod'][key]['blur'] if key in ctx['prod'] else pic(ctx, key)


def plain(t):
    return t.replace('**', '')


def rich(s, x, y, w, h, content, size=12.5, color=None, bold_color=None, align='l', line_sp=1.25, space_after=6,
         nm=None, anchor='t', bullets=False):
    color = color or T['text']
    bold_color = bold_color or T['accent']
    paras = []
    for para in content.split('\n'):
        runs = []
        for i, part in enumerate(re.split(r'\*\*', para)):
            if part:
                runs.append((part, {'font': F('bold'), 'bold': BOLD(), 'color': bold_color}) if i % 2 else (part, {}))
        paras.append(runs or [''])
    return K.text(s, x, y, w, h, paras, size=size, font=F('body'), color=color, align=align, line_sp=line_sp,
                  space_after=space_after, nm=nm, anchor=anchor, bullets=bullets, bullet_color=T['accent'])


def _effects(tb, xml):
    for r in tb._element.iter(K.qn('a:rPr')):
        old = r.find(K.qn('a:effectLst'))
        if old is not None:
            r.remove(old)
        eff = K.X(xml)
        first = r.find(K.qn('a:latin'))
        if first is not None:
            first.addprevious(eff)


def extrude(tb, depth=0.07, color='A9AEB2'):
    """Pseudo-3D: white→grey gradient face + hard offset shadow (stays editable text)."""
    for r in tb._element.iter(K.qn('a:rPr')):
        sf = r.find(K.qn('a:solidFill'))
        g = K.X('<a:gradFill rotWithShape="1"><a:gsLst><a:gs pos="0"><a:srgbClr val="FFFFFF"/></a:gs>'
                '<a:gs pos="60000"><a:srgbClr val="F1F2F2"/></a:gs><a:gs pos="100000"><a:srgbClr val="C9CDD0"/></a:gs>'
                '</a:gsLst><a:lin ang="5400000" scaled="0"/></a:gradFill>')
        if sf is not None:
            sf.addprevious(g)
            r.remove(sf)
    _effects(tb, f'<a:effectLst><a:outerShdw blurRad="0" dist="{int(depth * 914400)}" dir="2700000" algn="ctr" '
                 f'rotWithShape="0"><a:srgbClr val="{color}"/></a:outerShdw></a:effectLst>')


def soft_shadow(tb, alpha=35, blur=0.12, dist=0.05, color='000000'):
    _effects(tb, f'<a:effectLst><a:outerShdw blurRad="{int(blur * 914400)}" dist="{int(dist * 914400)}" '
                 f'dir="5400000" algn="ctr" rotWithShape="0"><a:srgbClr val="{color}"><a:alpha val="{alpha * 1000}"/>'
                 f'</a:srgbClr></a:outerShdw></a:effectLst>')


def stretch(s, path, x, y, w, h, nm=None):
    """Picture stretched to the box (no cover-cropping) — for soft glows / washes."""
    p = s.shapes.add_picture(path, K.E(x), K.E(y), K.E(w), K.E(h))
    if nm:
        p.name = nm
    return p


def product(s, ctx, key, cx, bottom, h, nm=None, rot=0, blurred=False, w=None):
    path = blur(ctx, key) if blurred else pic(ctx, key)
    from PIL import Image
    iw, ih = Image.open(path).size
    w = w or h * iw / ih
    if w * ih / iw < h:
        h = w * ih / iw
    p = K.picture(s, path, cx - w / 2, bottom - h, w, h, nm=nm, rot=rot)
    return p


def decor(s, ctx, spots, prefix='!!d'):
    """Flying decorations: spots = [(art key, cx, cy, size, rot, blurred)]; named so Morph floats them."""
    out = []
    for i, (key, cx, cy, d, rot, bl) in enumerate(spots):
        path = ctx['art'][key + ('_blur' if bl else '')]
        from PIL import Image
        iw, ih = Image.open(path).size
        w, h = (d, d * ih / iw) if iw >= ih else (d * iw / ih, d)
        out.append(K.picture(s, path, cx - w / 2, cy - h / 2, w, h, rot=rot, nm=f'{prefix}{i}'))
    return out


def decor_set(ctx, idx, n=5, avoid=None):
    """Deterministic pseudo-random decoration spots for a slide index (theme-specific art keys).
    avoid = (y0, y1): keep decorations out of this horizontal band (e.g. a giant word)."""
    import random
    rnd = random.Random(idx * 7 + 3)
    keys = ctx['decor']
    spots = []
    for i in range(n):
        side = i % 2
        cx = rnd.uniform(0.3, 2.6) if side == 0 else rnd.uniform(10.6, 13.1)
        cy = rnd.uniform(1.3, 6.9)
        if avoid and avoid[0] - 0.5 < cy < avoid[1] + 0.5:
            cy = avoid[0] - 0.55 if (i // 2) % 2 == 0 else avoid[1] + 0.55
        spots.append((keys[i % len(keys)], cx, cy, rnd.uniform(0.45, 1.05), rnd.uniform(-60, 60), i >= 2))
    return spots


def background(s, ctx, photo=None, focus=(0.5, 0.5), wash=True):
    """Theme background: paper (M07) or white (M06); optional photo with a wash at the top for the nav bar."""
    out = []
    if photo:
        out.append(K.picture(s, pic(ctx, photo), 0, 0, SW, SH, focus=tuple(focus), nm='!!bg'))
        if wash:
            ov = K.rect(s, 0, 0, SW, SH, nm='!!wash')
            K.grad_fill(ov, [(0, T['bg'], 92), (16, T['bg'], 55), (40, T['bg'], 0), (100, T['bg'], 0)], angle=90)
            out.append(ov)
    elif T['paper']:
        out.append(K.picture(s, ctx['art']['paper'], 0, 0, SW, SH, nm='!!bg'))
    else:
        out.append(K.picture(s, ctx['art']['white'], 0, 0, SW, SH, nm='!!bg'))
    return out


def nav(s, ctx, active=0, dark=False):
    """Website navigation bar (logo, menu, icons) — same names on every slide so it stays still under Morph."""
    b = ctx['brand']
    col = 'FFFFFF' if dark else T['ink']
    out = []
    ic = K.icon(s, b.get('icon', 'sprout'), ML, 0.3, 0.34, T['accent'] if not dark else 'FFFFFF', nmx='!!nav_logo')
    out.append(ic)
    out.append(K.text(s, ML + 0.42, 0.25, 3.0, 0.44, b['name'].upper(), size=13, font=F('title'), color=T['title']
                      if not dark else 'FFFFFF', anchor='m', spacing=1.2, nm='!!nav_brand', wrap=False))
    menu = b.get('menu', [])
    x = 4.3
    pos = []
    for i, m in enumerate(menu):
        w = K.text_w(m, F('semi'), 10.5) + 0.15
        out.append(K.text(s, x, 0.27, w, 0.4, m, size=10.5, font=F('semi'), color=col, anchor='m', nm=f'!!nav_m{i}',
                          wrap=False))
        pos.append((x, w))
        x += w + 0.45
    if menu:
        ax, aw = pos[min(active, len(pos) - 1)]
        out.append(K.rect(s, ax, 0.66, aw - 0.1, 0.035, fill=T['accent'], nm='!!nav_ind'))
    for i, nmi in enumerate(('search', 'user-round', 'shopping-bag')):
        out.append(K.icon(s, nmi, SW - ML - 0.28 - i * 0.48, 0.33, 0.27, col, nmx=f'!!nav_i{i}'))
    return out


def big_title(s, x, y, w, h, t, size, align='c', nm=None, color=None, anchor='m'):
    tb = K.text(s, x, y, w, h, t, size=size, font=F('title'), color=color or T['title'], align=align, anchor=anchor,
                nm=nm, wrap=False, line_sp=0.9)
    return tb


def fit(t, role, size, w, lo=12):
    while size > lo and K.text_w(t, F(role), size) > w:
        size -= 1
    return size


def script(s, x, y, w, h, t, size, align='c', nm=None, color=None):
    return K.text(s, x, y, w, h, t, size=size, font=F('script'), color=color or T['title'], align=align,
                  anchor='m', nm=nm, wrap=False)


def heading(s, title, sub='', y=1.05, align='c', x=ML, w=SW - 2 * ML, size=34):
    """Bold title + script word overlapping it (TH / Highlands headline style)."""
    sz = fit(title.upper(), 'title', size, w - 0.4, 20)
    t = big_title(s, x, y, w, 0.8, title.upper(), sz, align=align, nm='!!h_title')
    out = [t]
    if sub:
        ss = int(sz * 0.95)
        off = 0.0 if align == 'c' else 0.35
        out.append(script(s, x + off, y + 0.42, w, 0.75, sub, ss, align=align, nm='!!h_script', color=T['ink']
                          if T['title_style'] == 'extrude' else T['ink']))
    return out


def card(s, x, y, w, h, radius=0.16, fill=None, alpha=None, nm=None, line=None):
    c = K.rect(s, x, y, w, h, fill=fill or T['card'], radius=radius, alpha=alpha, nm=nm, line=line, lw=0.75)
    K.shadow(c, blur=0.3, dist=0.08, alpha=16, color='3B2A20', angle=90)
    K._ensure_ln_before_effects(c)
    return c


def page_no(s, ctx):
    return K.text(s, SW - ML - 0.8, SH - 0.4, 0.8, 0.25, f'{ctx["n"]:02d}', size=9, font=F('semi'), color=T['muted'],
                  align='r')


# ================================================================== 1. cover
def cover(prs, ctx, c):
    s = _new(prs, ctx)
    bgs = background(s, ctx, c.get('photo'), c.get('focus', (0.5, 0.6)), wash=False)
    if T['paper'] and not c.get('photo'):
        bgs.append(K.picture(s, ctx['art']['grass'], 0, SH - 2.1, SW, 2.1, nm='!!grass'))
    word = c['word'].upper()
    size = 200
    while K.text_w(word, F('title'), size) > SW - 0.6 and size > 80:
        size -= 4
    wy = c.get('word_y', 1.15)
    wt = big_title(s, 0.3, wy, SW - 0.6, size / 72 * 1.05, word, size, nm='!!word')
    if T['title_style'] == 'extrude':
        extrude(wt, depth=0.09)
    else:
        soft_shadow(wt, alpha=25, blur=0.08, dist=0.04, color='5A0A14')
    dec = decor(s, ctx, decor_set(ctx, 0, 6, avoid=(wy, wy + size / 72 * 1.05)))
    prods = []
    items = c['products']
    main = items[0]
    prods.append(product(s, ctx, main, SW / 2, c.get('bottom', 6.75), c.get('height', 4.9), nm=f'!!p_{main}',
                         rot=c.get('rot', -4)))
    for k, (key, dx, h, rot) in enumerate(zip(items[1:3], (-3.0, 3.1), (3.2, 3.0), (-10, 8))):
        prods.insert(0, product(s, ctx, key, SW / 2 + dx, c.get('bottom', 6.75) - 0.15, h, nm=f'!!p_{key}', rot=rot))
    sc = None
    if c.get('script'):
        ssz = int(size * 0.42)
        sc = script(s, SW / 2 + 0.6, wy + size / 72 * 0.72, SW / 2 - 0.9, size / 72 * 0.6, c['script'], ssz,
                    align='l', nm='!!word_script', color=T['title'])
        _effects(sc, '<a:effectLst><a:glow rad="101600"><a:srgbClr val="FFFFFF"><a:alpha val="85000"/></a:srgbClr>'
                     '</a:glow></a:effectLst>')
    kick = None
    if c.get('kicker'):
        kt = c['kicker'].upper()
        kw = K.text_w(kt, F('semi'), 11) + len(kt) * 2 / 72 + 0.5
        kick = K.rect(s, ML, 0.45, kw, 0.36, fill=T['accent'], radius=0.18, nm='!!kicker')
        K.shape_text(kick, kt, size=11, color='FFFFFF', font=F('semi'), spacing=2, margin=(0.1, 0.02))
    logo = K.text(s, SW - ML - 4, 0.42, 4, 0.42, ctx['brand']['name'].upper(), size=14, font=F('title'),
                  color=T['title'], align='r', anchor='m', spacing=1.5, nm='!!nav_brand')
    tag = None
    if c.get('tagline'):
        on_photo = bool(c.get('photo')) or T['paper']
        tag = K.text(s, ML, SH - 0.85, 6.5, 0.5, c['tagline'], size=13, font=F('semi'), color='FFFFFF'
                     if on_photo else T['ink'], anchor='m')
        if on_photo:
            soft_shadow(tag, alpha=75, blur=0.1, dist=0.03)
    K.anim(s, bgs[0], 'fade', 0.0, 0.8)
    K.anim(s, wt, 'zoom', 0.2, 1.0)
    if sc is not None:
        K.anim(s, sc, 'wipe_l', 0.9, 0.9)
    K.anims(s, prods, 'rise', 0.6, 0.15, 0.8)
    K.anims(s, dec, 'fade', 1.0, 0.08, 0.6)
    for x in (kick, logo, tag):
        if x is not None:
            K.anim(s, x, 'fade', 1.3, 0.6)
    K.notes(s, c.get('notes', 'Slide bìa.'))
    return s


# ================================================================== 2. intro
def intro(prs, ctx, c):
    s = _new(prs, ctx)
    background(s, ctx)
    hd = heading(s, c['title'], c.get('script', ''), y=1.0)
    tx = rich(s, 2.0, 2.45, SW - 4.0, 1.8, c['text'], size=13.5, align='c', line_sp=1.35)
    strip = []
    if c.get('photo'):
        p = K.picture(s, pic(ctx, c['photo']), 0, 4.55, SW, SH - 4.55, focus=tuple(c.get('focus', (0.5, 0.5))),
                      nm='!!strip')
        ov = K.rect(s, 0, 4.55, SW, 1.2)
        K.grad_fill(ov, [(0, T['bg'], 100), (100, T['bg'], 0)], angle=90)
        strip = [p, ov]
    elif T['paper']:
        strip = [K.picture(s, ctx['art']['grass'], 0, SH - 2.1, SW, 2.1, nm='!!grass')]
    m = K.picture(s, ctx['art']['mist'], -0.5, SH - 1.5, SW + 1, 1.7, nm='!!mist') if not T['paper'] else None
    dec = decor(s, ctx, decor_set(ctx, 1, 4))
    K.anims(s, strip, 'fade', 0.0, 0.0, 0.8)
    K.anims(s, hd, 'zoom', 0.3, 0.2, 0.7)
    K.anim(s, tx, 'fade', 0.8, 0.7)
    if m is not None:
        K.anim(s, m, 'fade', 0.2, 1.0)
    K.anims(s, dec, 'fade', 0.6, 0.1, 0.5)
    return s


# ================================================================== 3. story cards (sepia photos)
def story_cards(prs, ctx, c):
    s = _new(prs, ctx)
    background(s, ctx)
    nv = nav(s, ctx, c.get('nav', 0))
    cards = c['cards'][:3]
    n = len(cards)
    gap = 0.3
    h = (SH - 1.25 - 0.45 - gap * (n - 1)) / n
    out = []
    for i, cd in enumerate(cards):
        y = 1.2 + i * (h + gap)
        x = ML + (0.6 if i % 2 else 0)
        w = SW - 2 * ML - 0.6
        bx = card(s, x, y, w, h, radius=0.14)
        ph = K.picture(s, ctx['art']['sepia_' + cd['photo']], x + (w * 0.45 if i % 2 == 0 else 0), y, w * 0.55, h,
                       radius=0.14)
        fade = K.rect(s, x + (w * 0.45 if i % 2 == 0 else 0), y, w * 0.55, h, radius=0.14)
        K.grad_fill(fade, [(0, T['card'], 100), (45, T['card'], 40), (100, T['card'], 0)],
                    angle=0 if i % 2 == 0 else 180)
        tx0 = x + 0.35 if i % 2 == 0 else x + w * 0.45
        tw = w * 0.52
        title = cd['title'].upper()
        sz = fit(title, 'title', 30, tw - 0.2, 18)
        tt = big_title(s, tx0, y + 0.12, tw, 0.6, title, sz, align='l')
        sc = script(s, tx0 + K.text_w(title, F('title'), sz) * 0.55, y + 0.38, tw, 0.55, cd.get('script', ''), 26,
                    align='l', color=T['ink']) if cd.get('script') else None
        body = rich(s, tx0, y + 0.95, tw - 0.1, h - 1.05, cd['text'], size=11.5, line_sp=1.2)
        out.append([bx, ph, fade, tt] + ([sc] if sc is not None else []) + [body])
    K.anims(s, nv, 'fade', 0.0, 0.0, 0.4)
    K.anims(s, out, 'fly_l', 0.2, 0.25, 0.7)
    page_no(s, ctx)
    return s


# ================================================================== 4. story with big product
def story_product(prs, ctx, c):
    s = _new(prs, ctx)
    background(s, ctx)
    nv = nav(s, ctx, c.get('nav', 0))
    right = c.get('side', 'right') == 'right'           # side of the product
    dec = decor(s, ctx, decor_set(ctx, ctx['n'], 5))
    px = 9.4 if right else 3.9
    p = product(s, ctx, c['product'], px, 7.05, c.get('height', 5.6), nm=f'!!p_{c["product"]}', rot=c.get('rot', -6))
    p2 = product(s, ctx, c['second'], px + (2.4 if right else -2.4), 6.6, 2.4, nm=f'!!p_{c["second"]}', rot=10) \
        if c.get('second') else None
    tx0 = ML + 0.2 if right else 6.9
    tw = 5.9
    title = c['title'].upper()
    sz = fit(title, 'title', 54, tw, 28)
    tt = big_title(s, tx0, 2.0, tw, 1.0, title, sz, align='l', nm='!!h_title')
    sc = script(s, tx0 + 0.5, 2.65, tw, 0.8, c.get('script', ''), 34, align='l', nm='!!h_script', color=T['ink']) \
        if c.get('script') else None
    body = rich(s, tx0, 3.6, tw, 2.6, c['text'], size=13, line_sp=1.35)
    K.anims(s, nv, 'fade', 0.0, 0.0, 0.4)
    K.anim(s, p, 'fly_r' if right else 'fly_l', 0.2, 0.9)
    if p2 is not None:
        K.anim(s, p2, 'rise', 0.5, 0.7)
    K.anims(s, dec, 'fade', 0.4, 0.08, 0.6)
    K.anim(s, tt, 'wipe_l', 0.5, 0.6)
    if sc is not None:
        K.anim(s, sc, 'fade', 0.9, 0.6)
    K.anim(s, body, 'fade', 1.0, 0.7)
    page_no(s, ctx)
    return s


# ================================================================== 5. values row
def values_row(prs, ctx, c):
    s = _new(prs, ctx)
    background(s, ctx)
    nv = nav(s, ctx, c.get('nav', 0))
    hd = heading(s, c['title'], c.get('script', ''), y=1.05)
    tx = rich(s, 1.6, 2.35, SW - 3.2, 1.2, c['text'], size=13, align='c', line_sp=1.3)
    items = c['items'][:4]
    n = len(items)
    step = (SW - 2 * ML) / n
    out = []
    for i, it in enumerate(items):
        cx = ML + step * (i + 0.5)
        p = product(s, ctx, it['product'], cx, 6.55, 2.7, nm=f'!!p_{it["product"]}', rot=(-5, 4, -3, 6)[i % 4])
        lab = K.text(s, cx - step / 2 + 0.1, 6.6, step - 0.2, 0.36, it['label'].upper(), size=13, font=F('title'),
                     color=T['title'], align='c')
        sub = K.text(s, cx - step / 2 + 0.2, 6.95, step - 0.4, 0.35, it.get('text', ''), size=10, font=F('body'),
                     color=T['text'], align='c')
        out.append([p, lab, sub])
    K.anims(s, nv, 'fade', 0.0, 0.0, 0.4)
    K.anims(s, hd, 'zoom', 0.2, 0.15, 0.6)
    K.anim(s, tx, 'fade', 0.6, 0.6)
    K.anims(s, out, 'rise', 0.9, 0.18, 0.6)
    page_no(s, ctx)
    return s


# ================================================================== 6. line-up cards
def _split_word(word):
    if ' ' in word.strip():
        parts = word.split()
        k = len(parts) // 2
        return ' '.join(parts[:k]), ' '.join(parts[k:])
    k = len(word) // 2
    return word[:k], word[k:]


def prod_w(ctx, key, h):
    from PIL import Image
    iw, ih = Image.open(pic(ctx, key)).size
    return h * iw / ih


def word_with_product(s, ctx, word, key, cx, cy, size, ph, gap=None, nm_word=None, nm_prod=None, rot=-6,
                      style=None):
    """A big word with the product in front of it (style 'front') or in a gap between its halves ('gap')."""
    style = style or T['word_style']
    out = []
    if style == 'gap':
        a, b = _split_word(word.upper())
        gap = gap if gap is not None else prod_w(ctx, key, ph) * 1.02 + 0.3
        wa, wb = K.text_w(a, F('title'), size), K.text_w(b, F('title'), size)
        tot = wa + gap + wb
        x0 = cx - tot / 2
        h = size / 72 * 1.25
        out.append(big_title(s, x0 - 0.1, cy - h / 2, wa + 0.2, h, a, size, align='r', nm=nm_word and nm_word + 'a'))
        out.append(big_title(s, x0 + wa + gap - 0.1, cy - h / 2, wb + 0.2, h, b, size, align='l',
                             nm=nm_word and nm_word + 'b'))
        out.append(product(s, ctx, key, x0 + wa + gap / 2, cy + ph * 0.5, ph, nm=nm_prod, rot=rot))
    else:
        h = size / 72 * 1.25
        w = K.text_w(word.upper(), F('title'), size) + 0.4
        out.append(big_title(s, cx - w / 2, cy - h / 2, w, h, word.upper(), size, nm=nm_word))
        out.append(product(s, ctx, key, cx, cy + ph * 0.5, ph, nm=nm_prod, rot=rot))
    return out


def lineup(prs, ctx, c):
    s = _new(prs, ctx)
    background(s, ctx)
    nv = nav(s, ctx, c.get('nav', 0))
    hd = heading(s, c['title'], c.get('script', ''), y=1.0)
    items = c['items'][:3]
    n = len(items)
    gap = 0.4
    w = (SW - 2 * ML - gap * (n - 1)) / n
    out = []
    for i, it in enumerate(items):
        x = ML + i * (w + gap)
        y = 2.9
        cd = card(s, x, y, w, 2.6, alpha=None)
        ph = 2.05
        a_, b_ = _split_word(it['word'].upper())
        room = w - 0.45 - prod_w(ctx, it['product'], ph) * 1.02 - 0.3
        size = 46
        while size > 18 and K.text_w(a_, F('title'), size) + K.text_w(b_, F('title'), size) > room:
            size -= 1
        grp = word_with_product(s, ctx, it['word'], it['product'], x + w / 2, y + 1.3, size, ph,
                                nm_prod=f'!!p_{it["product"]}', style='gap', rot=-4)
        cap = K.text(s, x, y + 2.75, w, 0.4, it.get('caption', it['word']), size=13, font=F('semi'),
                     color=T['ink'], align='c')
        out.append([cd] + grp + [cap])
    K.anims(s, nv, 'fade', 0.0, 0.0, 0.4)
    K.anims(s, hd, 'zoom', 0.2, 0.15, 0.6)
    K.anims(s, out, 'rise', 0.6, 0.2, 0.6)
    page_no(s, ctx)
    return s


# ================================================================== 7. category word
def category(prs, ctx, c):
    s = _new(prs, ctx)
    background(s, ctx)
    nv = nav(s, ctx, c.get('nav', 0))
    word = c['word'].upper()
    size = 190
    lim = SW - (3.4 if T['word_style'] == 'gap' else 1.2)
    while K.text_w(word, F('title'), size) > lim and size > 70:
        size -= 4
    grp = word_with_product(s, ctx, c['word'], c['product'], SW / 2, 3.95, size, 4.6, nm_word='!!catw',
                            nm_prod=f'!!p_{c["product"]}', rot=c.get('rot', -8))
    sc = None
    if c.get('script'):
        sc = script(s, SW / 2 + 1.75, 4.75, 4.6, 1.2, c['script'], 60, align='l', nm='!!cat_script', color=T['ink'])
        soft_shadow(sc, alpha=80, blur=0.08, dist=0.0, color='FFFFFF')
    cap = K.text(s, ML, SH - 0.9, SW - 2 * ML, 0.4, c.get('caption', ''), size=13, font=F('semi'), color=T['ink'],
                 align='c') if c.get('caption') else None
    hw = size / 72 * 0.62
    dec = decor(s, ctx, decor_set(ctx, ctx['n'], 4, avoid=(3.95 - hw, 3.95 + hw)))
    K.anims(s, nv, 'fade', 0.0, 0.0, 0.4)
    K.anims(s, grp[:-1], 'zoom', 0.1, 0.1, 0.8)
    K.anim(s, grp[-1], 'rise', 0.5, 0.8)
    if sc is not None:
        K.anim(s, sc, 'wipe_l', 1.0, 0.8)
    if cap is not None:
        K.anim(s, cap, 'fade', 1.2, 0.5)
    K.anims(s, dec, 'fade', 0.6, 0.08, 0.5)
    page_no(s, ctx)
    return s


# ================================================================== 8. carousel (one slide per product)
SLOTS = [  # (cx, bottom, height, rot) for the non-hero products, nearest first
    (2.35, 3.1, 1.55, -14), (11.15, 4.3, 1.9, 12), (1.2, 5.6, 1.05, 18), (12.4, 2.4, 1.0, -20), (9.6, 2.2, 0.8, 8)]


def carousel(prs, ctx, c):
    items = c['items']
    slides = []
    for a, it in enumerate(items):
        s = _new(prs, ctx, tr=c.get('tr', 1.1))
        bg = background(s, ctx, it.get('photo', c.get('photo')), it.get('focus', c.get('focus', (0.5, 0.55))))
        mist = K.picture(s, ctx['art']['mist'], -0.5, SH - 2.6, SW + 1, 2.8, nm='!!mist') if c.get('mist', True) \
            else None
        stretch(s, K.glow_png(T['bg'], 92, falloff=1.1), -2.2, 3.2, 10.4, 5.6, nm='!!c_wash')
        nv = nav(s, ctx, c.get('nav', 0))
        others = [k for k in range(len(items)) if k != a]
        others.sort(key=lambda k: (min(abs(k - a), len(items) - abs(k - a)), k))
        side = []
        for slot, k in zip(SLOTS, others):
            cx, bot, h, rot = slot
            side.append(product(s, ctx, items[k]['product'], cx, bot, h, nm=f'!!p_{items[k]["product"]}', rot=rot,
                                blurred=True))
        hero = product(s, ctx, it['product'], SW / 2 + 0.4, 6.55, c.get('height', 5.0), nm=f'!!p_{it["product"]}',
                       rot=it.get('rot', -5))
        # name block
        name = it['name'].upper()
        nsz = fit(name, 'title', 44, 5.6, 24)
        nt = K.text(s, ML, 4.75, 6.0, 0.8, name, size=nsz, font=F('title'), color='FFFFFF' if c.get('dark_text') is
                    False else T['title'], anchor='b', nm='!!c_name', wrap=False)
        soft_shadow(nt, alpha=30, blur=0.1, dist=0.03)
        sc = script(s, ML + 0.6, 5.3, 6.0, 0.8, it.get('script', ''), 36, align='l', nm='!!c_script',
                    color=T['ink']) if it.get('script') else None
        if sc is not None:
            soft_shadow(sc, alpha=30, blur=0.1, dist=0.03, color='FFFFFF')
        ds = rich(s, ML, 6.15, 5.2, 0.9, it.get('text', ''), size=10.5, color=T['ink'], line_sp=1.15, nm='!!c_text')
        # thumbnails
        n = len(items)
        tw = 0.5
        x0 = SW - ML - n * (tw + 0.12) + 0.12
        thumbs = []
        for k, ti in enumerate(items):
            thumbs.append(product(s, ctx, ti['product'], x0 + k * (tw + 0.12) + tw / 2, 6.95, 0.62, nm=f'!!t_{k}',
                                  w=tw))
        ind = K.rect(s, x0 + a * (tw + 0.12), 7.05, tw, 0.05, fill=T['accent'], nm='!!t_ind')
        if a == 0:
            K.anims(s, nv, 'fade', 0.0, 0.0, 0.4)
            K.anim(s, hero, 'rise', 0.2, 0.8)
            K.anims(s, side, 'fade', 0.5, 0.1, 0.6)
            K.anims(s, thumbs + [ind], 'fade', 0.9, 0.04, 0.4)
        K.anim(s, nt, 'wipe_u', 0.3, 0.6)
        if sc is not None:
            K.anim(s, sc, 'wipe_l', 0.6, 0.7)
        K.anim(s, ds, 'fade', 0.8, 0.6)
        _ = (bg, mist)
        page_no(s, ctx)
        slides.append(s)
    return slides


# ================================================================== 9. features around the product
def features(prs, ctx, c):
    s = _new(prs, ctx)
    background(s, ctx)
    nv = nav(s, ctx, c.get('nav', 1))
    hd = heading(s, c['title'], c.get('script', ''), y=0.95, size=32)
    glow = K.glow_img(s, SW / 2, 4.6, 4.6, T['accent2'], 45)
    p = product(s, ctx, c['product'], SW / 2, 7.0, 4.6, nm=f'!!p_{c["product"]}', rot=-4)
    feats = c['items'][:4]
    pos = [(ML, 2.35, 'l'), (ML, 4.75, 'l'), (SW - ML - 3.9, 2.35, 'r'), (SW - ML - 3.9, 4.75, 'r')]
    out = []
    for (x, y, side), (ic, t, txt) in zip(pos, feats):
        bd = K.icon_badge(s, ic, x + (0.38 if side == 'l' else 3.52), y + 0.38, 0.72, T['accent'], 'FFFFFF',
                          icon_scale=0.5)
        tx = x + (0.9 if side == 'l' else 0)
        tt = K.text(s, tx, y + 0.02, 3.0, 0.42, t, size=15, font=F('bold'), bold=BOLD(), color=T['ink'], align=side, anchor='m')
        bt = K.text(s, tx, y + 0.45, 3.0, 1.2, txt, size=11, font=F('body'), color=T['text'], align=side,
                    line_sp=1.2)
        lx0 = x + 3.95 if side == 'l' else x - 0.05
        lx1 = SW / 2 - 1.3 if side == 'l' else SW / 2 + 1.3
        ln = K.line(s, lx0, y + 0.38, lx1, y + 0.38 + (0.4 if y < 3 else -0.3), T['accent'], 1.0, dash='dash')
        dot = K.oval(s, lx1 - 0.06, y + 0.32 + (0.4 if y < 3 else -0.3), 0.12, fill=T['accent'])
        out.append(bd + [tt, bt, ln, dot])
    K.anims(s, nv, 'fade', 0.0, 0.0, 0.4)
    K.anims(s, hd, 'zoom', 0.1, 0.15, 0.6)
    K.anim(s, glow, 'fade', 0.3, 0.8)
    K.anim(s, p, 'rise', 0.4, 0.8)
    K.anims(s, out, 'fade', 0.9, 0.2, 0.5)
    page_no(s, ctx)
    return s


# ================================================================== 10. spec table
def specs(prs, ctx, c):
    s = _new(prs, ctx)
    background(s, ctx)
    nv = nav(s, ctx, c.get('nav', 1))
    p = product(s, ctx, c['product'], 3.35, 6.95, 5.4, nm=f'!!p_{c["product"]}', rot=-5)
    x0 = 6.3
    w = SW - ML - x0
    title = c['title'].upper()
    tt = big_title(s, x0, 1.0, w, 0.75, title, fit(title, 'title', 36, w, 20), align='l', nm='!!h_title')
    sc = script(s, x0 + 0.4, 1.5, w, 0.7, c.get('script', ''), 30, align='l', nm='!!h_script', color=T['ink']) \
        if c.get('script') else None
    rows = c['rows']
    rh = min(0.5, 3.9 / len(rows))
    cw = c.get('col_w') or [2.6] + [1.6] * (len(rows[0]) - 1)
    k = w / sum(cw)
    gf = K.table(s, x0, 2.4, w, rh * len(rows), rows, col_w=[v * k for v in cw], header_fill=T['accent'],
                 header_color='FFFFFF', body_fill=T['card'], body_fill2=T['soft'], body_color=T['ink'],
                 line_color=T['line'], size=11.5, header_size=12, row_h=rh)
    note = rich(s, x0, 2.5 + rh * len(rows), w, 0.9, c.get('note', ''), size=10.5, color=T['muted'], line_sp=1.2)
    K.anims(s, nv, 'fade', 0.0, 0.0, 0.4)
    K.anim(s, p, 'fly_l', 0.2, 0.8)
    K.anim(s, tt, 'wipe_l', 0.4, 0.6)
    if sc is not None:
        K.anim(s, sc, 'fade', 0.7, 0.5)
    K.anim(s, gf, 'wipe_d', 0.8, 0.9)
    K.anim(s, note, 'fade', 1.5, 0.5)
    page_no(s, ctx)
    return s


# ================================================================== 11. usage steps
def usage(prs, ctx, c):
    s = _new(prs, ctx)
    background(s, ctx)
    nv = nav(s, ctx, c.get('nav', 2))
    hd = heading(s, c['title'], c.get('script', ''), y=0.95, size=32)
    steps = c['steps'][:5]
    n = len(steps)
    x0 = ML + (2.6 if c.get('product') else 0)
    w = SW - ML - x0
    step = w / n
    p = product(s, ctx, c['product'], ML + 1.25, 6.9, 4.2, nm=f'!!p_{c["product"]}', rot=-6) if c.get('product') \
        else None
    rail = K.line(s, x0 + step / 2, 3.15, x0 + w - step / 2, 3.15, T['accent'], 2, dash='dash')
    out = []
    for i, (ic, t, txt) in enumerate(steps):
        cx = x0 + step * (i + 0.5)
        circ = K.oval(s, cx - 0.55, 2.6, 1.1, fill=T['card'], line=T['accent'], lw=2)
        K.shadow(circ, blur=0.2, dist=0.05, alpha=18, color='3B2A20')
        K._ensure_ln_before_effects(circ)
        icn = K.icon(s, ic, cx - 0.28, 2.87, 0.56, T['accent'])
        num = K.oval(s, cx + 0.22, 2.5, 0.38, fill=T['title'])
        K.shape_text(num, str(i + 1), size=11, color='FFFFFF', font=F('bold'), bold=BOLD(), margin=(0, 0))
        tt = K.text(s, cx - step / 2 + 0.08, 3.9, step - 0.16, 0.6, t, size=14, font=F('bold'), bold=BOLD(), color=T['ink'],
                    align='c', line_sp=1.05)
        bt = K.text(s, cx - step / 2 + 0.1, 4.55, step - 0.2, 2.0, txt, size=11, font=F('body'), color=T['text'],
                    align='c', line_sp=1.2)
        out.append([circ, icn, num, tt, bt])
    K.anims(s, nv, 'fade', 0.0, 0.0, 0.4)
    K.anims(s, hd, 'zoom', 0.1, 0.15, 0.6)
    if p is not None:
        K.anim(s, p, 'fly_l', 0.3, 0.8)
    K.anim(s, rail, 'wipe_l', 0.6, 0.8)
    K.anims(s, out, 'rise', 0.8, 0.2, 0.5)
    page_no(s, ctx)
    return s


# ================================================================== 12. results (chart + KPIs + photo)
def results(prs, ctx, c):
    s = _new(prs, ctx)
    background(s, ctx)
    nv = nav(s, ctx, c.get('nav', 2))
    hd = heading(s, c['title'], c.get('script', ''), y=0.95, size=32, align='l', x=ML, w=8.0)
    ch = c['chart']
    cd = card(s, ML, 2.25, 7.3, 4.6)
    gf = K.chart_bar(s, ML + 0.25, 2.45, 6.8, 3.85, ch['cats'], [(n, v) for n, v in ch['series']],
                     colors=T['chart'], gap=70, overlap=-10, labels=True, num_fmt=ch.get('fmt', '0'), legend=True,
                     font_size=10.5, label_color=T['ink'], max_v=ch.get('max'))
    gch = gf.chart
    gch.font.name = F('body')
    for ax in (gch.category_axis, gch.value_axis):
        ax.tick_labels.font.name = F('body')
    note = K.text(s, ML + 0.3, 6.35, 6.8, 0.4, c.get('note', ''), size=9.5, font=F('body'), italic=True,
                  color=T['muted'], anchor='m')
    kp = []
    x0 = 8.25
    w = SW - ML - x0
    for i, (v, u, t) in enumerate(c.get('kpis', [])[:3]):
        y = 2.25 + i * 1.12
        b = card(s, x0, y, w, 0.98)
        vt = K.text(s, x0 + 0.25, y + 0.05, 2.2, 0.88, [[(v, {}), (u, {'size': 15})]], size=32, font=F('title'),
                    color=T['title'], anchor='m')
        lt = K.text(s, x0 + 2.2, y + 0.08, w - 2.35, 0.82, t, size=10.5, font=F('body'), color=T['text'],
                    anchor='m', line_sp=1.1)
        kp.append([b, vt, lt])
    ph = K.picture(s, pic(ctx, c['photo']), x0, 5.65, w, 1.2, radius=0.14, focus=tuple(c.get('focus', (0.5, 0.5)))) \
        if c.get('photo') else None
    K.anims(s, nv, 'fade', 0.0, 0.0, 0.4)
    K.anims(s, hd, 'wipe_l', 0.1, 0.15, 0.6)
    K.anim(s, cd, 'fade', 0.3, 0.5)
    K.anim(s, gf, 'wipe_u', 0.5, 0.9)
    K.anim(s, note, 'fade', 1.2, 0.4)
    K.anims(s, kp, 'fly_r', 0.6, 0.2, 0.5)
    if ph is not None:
        K.anim(s, ph, 'zoom', 1.3, 0.6)
    page_no(s, ctx)
    return s


# ================================================================== 13. browser-window cards
def window_cards(prs, ctx, c):
    s = _new(prs, ctx)
    background(s, ctx)
    nv = nav(s, ctx, c.get('nav', 0))
    items = c['items'][:3]
    out = []
    hd = heading(s, c.get('title', ''), c.get('script', ''), y=0.95, align='l', x=ML, w=6.0, size=30)
    geo = [(3.25, 2.05, 6.3, 4.75), (9.95, 2.45, 2.85, 2.5), (0.6, 3.75, 2.35, 2.35)]
    for i, (it, (x, y, w, h)) in enumerate(zip(items, geo)):
        win = card(s, x, y, w, h, radius=0.12)
        bar = K.rect(s, x, y, w, 0.32, fill=T['soft'], shp=MSO_SHAPE.ROUND_2_SAME_RECTANGLE)
        bar.adjustments[0] = 0.12 / 0.32
        dots = [K.oval(s, x + 0.15 + k * 0.17, y + 0.11, 0.1, fill=col) for k, col in
                enumerate(('E86A5B', 'F2C14E', '7CC26B'))]
        ph = K.picture(s, pic(ctx, it['photo']), x + 0.12, y + 0.42, w - 0.24, h - (1.05 if i == 0 else 0.54),
                       radius=0.08, focus=tuple(it.get('focus', (0.5, 0.5))))
        grp = [win, bar] + dots + [ph]
        if i == 0:
            nm = K.text(s, x + 0.25, y + h - 0.6, w - 0.5, 0.5, it['name'].upper(), size=18, font=F('title'),
                        color=T['title'], anchor='m')
            grp.append(nm)
        out.append(grp)
    tx = rich(s, 9.95, 5.15, SW - ML - 9.95, 1.7, c.get('text', ''), size=11.5, line_sp=1.25)
    K.anims(s, nv, 'fade', 0.0, 0.0, 0.4)
    K.anims(s, hd, 'wipe_l', 0.1, 0.1, 0.5)
    K.anims(s, out, 'zoom', 0.3, 0.25, 0.6)
    K.anim(s, tx, 'fade', 1.0, 0.6)
    page_no(s, ctx)
    return s


# ================================================================== 14. conclusion
def conclusion(prs, ctx, c):
    s = _new(prs, ctx)
    background(s, ctx)
    nv = nav(s, ctx, c.get('nav', 0))
    hd = heading(s, c['title'], c.get('script', ''), y=1.15)
    tx = rich(s, 1.8, 2.55, SW - 3.6, 1.5, c['text'], size=13.5, align='c', line_sp=1.35)
    items = c.get('products', [])[:5]
    n = len(items)
    out = []
    if n:
        step = min(2.3, (SW - 2 * ML) / n)
        x0 = SW / 2 - step * n / 2
        for i, k in enumerate(items):
            out.append(product(s, ctx, k, x0 + step * (i + 0.5), 6.85, 2.5, nm=f'!!p_{k}', rot=(-6, 4, -3, 7, -5)[i]))
    dec = decor(s, ctx, decor_set(ctx, ctx['n'], 4))
    K.anims(s, nv, 'fade', 0.0, 0.0, 0.4)
    K.anims(s, hd, 'zoom', 0.1, 0.15, 0.6)
    K.anim(s, tx, 'fade', 0.5, 0.7)
    K.anims(s, out, 'rise', 0.8, 0.12, 0.6)
    K.anims(s, dec, 'fade', 0.6, 0.1, 0.5)
    page_no(s, ctx)
    return s


# ================================================================== 15. contact / thanks
def contact(prs, ctx, c):
    s = _new(prs, ctx, tr=1.4)
    bgs = background(s, ctx, c.get('photo'), c.get('focus', (0.5, 0.6)), wash=False)
    if T['paper'] and not c.get('photo'):
        bgs.append(K.picture(s, ctx['art']['grass'], 0, SH - 2.1, SW, 2.1, nm='!!grass'))
    word = c.get('word', 'Xin cảm ơn').upper()
    size = 150
    while K.text_w(word, F('title'), size) > 7.6 and size > 60:
        size -= 4
    if c.get('photo'):
        lw = K.rect(s, 0, 0, 9.0, SH, nm='!!c_wash')
        K.grad_fill(lw, [(0, T['bg'], 88), (60, T['bg'], 60), (100, T['bg'], 0)], angle=0)
    wt = big_title(s, ML - 0.1, 0.9, 8.2, size / 72 * 1.1, word, size, align='l', nm='!!word')
    sc = script(s, ML + 1.0, 0.9 + size / 72 * 0.65, 7.0, 0.9, c.get('script', ''), 48, align='l',
                nm='!!word_script', color=T['title'] if T['title_style'] == 'flat' else T['accent']) \
        if c.get('script') else None
    p = product(s, ctx, c['product'], 10.2, 6.9, 5.4, nm=f'!!p_{c["product"]}', rot=-5)
    box = card(s, ML, 3.65, 6.6, 2.6, alpha=90)
    rows = []
    y = 3.85
    for ic, t in c.get('lines', [])[:5]:
        i_ = K.icon(s, ic, ML + 0.3, y + 0.05, 0.32, T['accent'])
        tx = K.text(s, ML + 0.8, y, 5.6, 0.42, t, size=13, font=F('semi'), color=T['ink'], anchor='m')
        rows.append([i_, tx])
        y += 0.46
    K.anim(s, bgs[0], 'fade', 0.0, 0.8)
    K.anim(s, wt, 'zoom', 0.2, 0.9)
    if sc is not None:
        K.anim(s, sc, 'wipe_l', 0.8, 0.8)
    K.anim(s, p, 'rise', 0.5, 0.9)
    K.anim(s, box, 'fade', 0.9, 0.5)
    K.anims(s, rows, 'wipe_l', 1.1, 0.12, 0.4)
    return s


# ================================================================== credits (hidden)
def credits(prs, ctx, c):
    s = _new(prs, ctx)
    s._element.set('show', '0')
    background(s, ctx)
    K.text(s, ML, 0.4, 11, 0.6, 'NGUỒN HÌNH ẢNH', size=26, font=F('title'), color=T['title'])
    K.text(s, ML, 1.05, 12, 0.5, 'Slide ẩn — không hiển thị khi trình chiếu. ' + c['intro'], size=10, font=F('body'),
           color=T['muted'])
    lines = c['lines']
    half = (len(lines) + 1) // 2
    for col, chunk in enumerate((lines[:half], lines[half:])):
        K.text(s, ML + col * 6.1, 1.75, 5.9, 5.4, '\n'.join(chunk), size=8, font=F('body'), color=T['text'],
               line_sp=1.05)
    return s
