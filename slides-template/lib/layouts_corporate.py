"""
Corporate layouts shared by M08 "Hồ sơ Doanh nghiệp" (Cotecons style: royal blue, glowing ring emblem, wide display
type) and M09 "Giới thiệu Dịch vụ" (Manulife style: fresh green, slanted bars and parallelogram photo panels).

Both themes use the same slide set; the theme switches palette, fonts, backgrounds and accents:
    style 'tech'    dark covers / sections on a glowing blue background, section pill top-left, rounded photos
    style 'chevron' green gradient covers, two-line titles with a slanted bar, parallelogram photos
Each layout: fn(prs, ctx, c) -> slide.
ctx = dict(n, img{key: path}, cut{key: path}, art{key: path}, brand{name, short, initial, tagline})
"""
import re

import pptkit as K
from pptx.enum.shapes import MSO_SHAPE

SW, SH, ML = K.SW, K.SH, 0.6

THEMES = {
    'doanhnghiep': dict(
        name='Hồ sơ Doanh nghiệp', style='tech',
        dark='061A5E', primary='1F4FD8', light='3FA9FF', pale='DCE7F7', ink='0B1F4D', text='4A5878', muted='8A97B5',
        bg='F4F7FB', card='FFFFFF', accent='FFC531', line='D5DEEE', chart=['1F4FD8', '3FA9FF', '0B1F4D', 'FFC531'],
        fonts=dict(title=('Lexend Exa', 'doanhnghiep/LexendExa-Bold.ttf', True),
                   title2=('Lexend Exa Medium', 'doanhnghiep/LexendExa-Medium.ttf', False),
                   body=('Be Vietnam Pro', 'doanhnghiep/BeVietnamPro-Regular.ttf', False),
                   semi=('Be Vietnam Pro SemiBold', 'doanhnghiep/BeVietnamPro-SemiBold.ttf', False),
                   bold=('Be Vietnam Pro', 'doanhnghiep/BeVietnamPro-Bold.ttf', True),
                   script=('Great Vibes', 'thiennhien/GreatVibes-Regular.ttf', False))),
    'dichvu': dict(
        name='Giới thiệu Dịch vụ', style='chevron',
        dark='0A5C36', primary='00A758', light='7FD6A3', pale='DFF3E8', ink='173A2A', text='4C5A52', muted='8FA096',
        bg='F5F8F6', card='FFFFFF', accent='1F3C99', line='D6E6DC', chart=['00A758', '7FD6A3', '0A5C36', '1F3C99'],
        fonts=dict(title=('Be Vietnam Pro ExtraBold', 'dichvu/BeVietnamPro-ExtraBold.ttf', False),
                   title2=('Be Vietnam Pro SemiBold', 'dichvu/BeVietnamPro-SemiBold.ttf', False),
                   body=('Be Vietnam Pro', 'dichvu/BeVietnamPro-Regular.ttf', False),
                   semi=('Be Vietnam Pro SemiBold', 'dichvu/BeVietnamPro-SemiBold.ttf', False),
                   bold=('Be Vietnam Pro', 'dichvu/BeVietnamPro-Bold.ttf', True),
                   script=('Great Vibes', 'thiennhien/GreatVibes-Regular.ttf', False))),
}
T = THEMES['doanhnghiep']


def F(role):
    return T['fonts'][role][0]


def B(role):
    return T['fonts'][role][2]


def register(key, colors=None):
    global T
    T = dict(THEMES[key])
    T.update(colors or {})
    K.THEMES['corp_' + key] = dict(
        name=T['name'], key='corp_' + key, bg_light=T['bg'], bg_dark=T['dark'], ink=T['ink'], text=T['text'],
        muted=T['muted'], soft=T['pale'], card=T['card'], card2=T['pale'], accent=T['primary'], accent_ink='FFFFFF',
        violet=T['light'], hot=T['accent'], up=T['primary'], line=T['line'], chart=T['chart'],
        blob_light=[(T['light'], 25)], blob_dark=[(T['light'], 25)], tint=T['dark'], grid_dark='2A3F7A',
        row_alt=T['pale'])
    K.set_theme('corp_' + key)
    for role in ('title', 'title2', 'body', 'semi', 'bold', 'script'):
        face, fname, bold = T['fonts'][role]
        if role == 'bold' and face in K.FONT_FILES and K.FONT_FILES[face] != fname:
            continue
        K.FONT_FILES[face] = fname
        if bold and role == 'title':          # Lexend Exa ships only Medium/Bold files
            K.BOLD_FACES.add(face)
    return T


TECH = lambda: T['style'] == 'tech'  # noqa: E731


# ------------------------------------------------------------------ helpers
def _new(prs, ctx, dark=False, tr=1.1):
    ctx['n'] += 1
    s = K.blank_slide(prs, dark)
    if ctx['n'] > 1:
        K.morph(s, tr)
    return s


def pic(ctx, key, prefer_cut=False):
    order = ('cut', 'img', 'art') if prefer_cut else ('img', 'art', 'cut')
    for k in order:
        if key in ctx[k]:
            return ctx[k][key]
    import content_io as CIO
    return CIO.photo_path(key)


def plain(t):
    return t.replace('**', '')


def rich(s, x, y, w, h, content, size=12.5, color=None, bold_color=None, align='l', line_sp=1.3, space_after=6,
         nm=None, anchor='t', bullets=False):
    color = color or T['text']
    bold_color = bold_color or T['primary']
    paras = []
    for para in content.split('\n'):
        runs = []
        for i, part in enumerate(re.split(r'\*\*', para)):
            if part:
                runs.append((part, {'font': F('bold'), 'bold': B('bold'), 'color': bold_color}) if i % 2 else (part, {}))
        paras.append(runs or [''])
    return K.text(s, x, y, w, h, paras, size=size, font=F('body'), color=color, align=align, line_sp=line_sp,
                  space_after=space_after, nm=nm, anchor=anchor, bullets=bullets, bullet_color=T['primary'])


def title_text(s, x, y, w, h, t, size, color=None, align='l', anchor='t', nm=None, role='title', wrap=True,
               line_sp=1.0, spacing=None):
    return K.text(s, x, y, w, h, t, size=size, font=F(role), bold=B(role), color=color or T['ink'], align=align,
                  anchor=anchor, nm=nm, wrap=wrap, line_sp=line_sp, spacing=spacing)


def fit(t, role, size, w, lo=12):
    while size > lo and K.text_w(t, F(role), size) > w:
        size -= 1
    return size


def glow_text(tb, color='FFFFFF', rad=0.08, alpha=40):
    for r in tb._element.iter(K.qn('a:rPr')):
        eff = K.X(f'<a:effectLst><a:glow rad="{int(rad * 914400)}"><a:srgbClr val="{color}"><a:alpha val="{alpha * 1000}"/>'
                  f'</a:srgbClr></a:glow></a:effectLst>')
        first = r.find(K.qn('a:latin'))
        if first is not None:
            first.addprevious(eff)


def bg_light(s, ctx):
    return K.picture(s, ctx['art']['light_bg'], 0, 0, SW, SH, nm='!!bg')


def bg_dark(s, ctx, key='dark_bg'):
    return K.picture(s, ctx['art'][key], 0, 0, SW, SH, nm='!!bg')


def logo(s, ctx, x=None, y=0.32, dark=False, size=0.42):
    """Brand mark top-right: emblem (tech) or slanted bars (chevron) + short name."""
    b = ctx['brand']
    name = b.get('short', b['name']).upper()
    nw = K.text_w(name, F('title'), 11) + len(name) * 1.5 / 72 + 0.1
    x = x if x is not None else SW - ML - nw - size - 0.1
    out = []
    if ctx['art'].get('logo'):
        out.append(K.picture(s, ctx['art']['logo'], x, y - 0.06, size + 0.12, size + 0.12, fit='contain',
                             nm='!!logo_mark'))
    elif TECH():
        out.append(K.picture(s, ctx['art']['emblem'], x, y - 0.05, size + 0.1, size + 0.1, nm='!!logo_mark'))
    else:
        for k in range(3):
            bar = K.rect(s, x + k * 0.13, y + 0.04, 0.09, size - 0.04, fill='FFFFFF' if dark else T['primary'],
                         shp=MSO_SHAPE.PARALLELOGRAM, nm=f'!!logo_bar{k}')
            bar.adjustments[0] = 0.45
            out.append(bar)
    out.append(K.text(s, x + size + 0.18, y, nw, size, name, size=11, font=F('title'), bold=B('title'),
                      color='FFFFFF' if dark else (T['primary'] if TECH() else T['ink']), anchor='m', spacing=1.5,
                      nm='!!logo_name', wrap=False))
    return out


def header(s, ctx, title, sub='', tag='', num=None, dark=False):
    """Content-slide header: tech = section pill + wide title; chevron = two-line title + slanted bar."""
    out = logo(s, ctx, dark=dark)
    ink = 'FFFFFF' if dark else T['ink']
    if TECH():
        if tag:
            tt = tag.upper()
            tw = K.text_w(tt, F('semi'), 9) + len(tt) * 1.0 / 72 + 0.75
            pill = K.rect(s, ML, 0.32, tw, 0.36, fill='FFFFFF', radius=0.18, line=T['line'], nm='!!tag')
            dot = K.oval(s, ML + 0.05, 0.37, 0.26, fill=T['primary'], nm='!!tag_dot')
            K.shape_text(dot, f'{num:02d}' if isinstance(num, int) else (num or ''), size=7, color='FFFFFF',
                         font=F('bold'), bold=B('bold'), margin=(0, 0))
            lb = K.text(s, ML + 0.38, 0.32, tw - 0.4, 0.36, tt, size=9, font=F('semi'), color=T['ink'], anchor='m',
                        spacing=1, nm='!!tag_text')
            out += [pill, dot, lb]
        t = title.upper()
        sz = fit(t, 'title', 26, 11.0, 16)
        out.append(title_text(s, ML, 0.85, 11.5, 0.6, t, sz, color=ink, anchor='m', nm='!!ttl', wrap=False))
        if sub:
            out.append(K.text(s, ML, 1.42, 10.5, 0.36, sub, size=12, font=F('body'), color='DCE6FF' if dark
                              else T['text'], nm='!!sub'))
    else:
        t = title.upper()
        sz = fit(t, 'title', 30, 8.5, 18)
        tt = title_text(s, ML, 0.4, 9.5, 0.62, t, sz, color=ink, anchor='b', nm='!!ttl', wrap=False)
        tw = K.text_w(t, F('title'), sz)
        bar = K.rect(s, ML + tw + 0.15, 0.62, 0.55, 0.16, fill=T['light'] if dark else T['primary'],
                     shp=MSO_SHAPE.PARALLELOGRAM, nm='!!bar')
        bar.adjustments[0] = 0.6
        out += [tt, bar]
        if sub:
            out.append(title_text(s, ML, 1.02, 9.5, 0.45, sub.upper(), fit(sub.upper(), 'title2', 20, 9.0, 12),
                                  color=T['light'] if dark else T['primary'], role='title2', nm='!!sub', wrap=False))
    out.append(K.text(s, SW - ML - 0.8, SH - 0.4, 0.8, 0.25, f'{ctx["n"]:02d}', size=9, font=F('semi'),
                      color='DCE6FF' if dark else T['muted'], align='r'))
    return out


def photo(s, ctx, key, x, y, w, h, focus=(0.5, 0.5), nm=None, shape=None):
    shape = shape or ('round' if TECH() else 'para')
    if shape == 'para':
        p = K.picture(s, pic(ctx, key), x, y, w, h, geom='parallelogram', focus=tuple(focus), nm=nm)
        g = p._element.spPr.find(K.qn('a:prstGeom'))
        av = g.find(K.qn('a:avLst'))
        gd = K.etree.SubElement(av, K.qn('a:gd'))
        gd.set('name', 'adj')
        gd.set('fmla', f'val {int(min(0.9, 0.9 * h / w) * 25000)}')
    else:
        p = K.picture(s, pic(ctx, key), x, y, w, h, radius=0.16, focus=tuple(focus), nm=nm)
    K.shadow(p, blur=0.3, dist=0.08, alpha=22, color='0B1F4D', angle=90)
    K._ensure_ln_before_effects(p)
    return p


def card(s, x, y, w, h, fill=None, radius=0.14, nm=None, line=None, alpha=None):
    c = K.rect(s, x, y, w, h, fill=fill or T['card'], radius=radius, line=line, lw=0.75, nm=nm, alpha=alpha)
    K.shadow(c, blur=0.28, dist=0.06, alpha=14, color='0B1F4D', angle=90)
    K._ensure_ln_before_effects(c)
    return c


def icon_badge(s, ic, cx, cy, d=0.62, fill=None, color='FFFFFF'):
    return K.icon_badge(s, ic, cx, cy, d, fill or T['primary'], color, icon_scale=0.52)


# ================================================================== 1. cover
def cover(prs, ctx, c):
    s = _new(prs, ctx, dark=True)
    b = ctx['brand']
    out = []
    if TECH():
        bg_dark(s, ctx, 'cover_bg')
        em = K.picture(s, ctx['art']['emblem'], 7.4, 0.9, 5.6, 5.6, nm='!!emblem')
        out.append(em)
        y = 2.0
    else:
        bg_dark(s, ctx, 'cover_bg')
        out.append(photo(s, ctx, c['photos'][0], 6.4, 0.0, 4.6, SH, focus=c.get('focus1', (0.5, 0.5)), nm='!!ph1'))
        if len(c['photos']) > 1:
            out.append(photo(s, ctx, c['photos'][1], 9.6, 0.0, 4.4, SH, focus=c.get('focus2', (0.5, 0.5)),
                             nm='!!ph2'))
        y = 2.1
    kt = c.get('kicker', '').upper()
    kick = None
    if kt:
        kw = K.text_w(kt, F('semi'), 10) + len(kt) * 1.5 / 72 + 0.5
        kick = K.rect(s, ML, y - 0.75, kw, 0.36, fill=T['light'] if TECH() else 'FFFFFF', radius=0.18, nm='!!kicker')
        K.shape_text(kick, kt, size=10, color=T['dark'], font=F('semi'), spacing=1.5, margin=(0.1, 0.02))
    l1 = K.text(s, ML, y, 6.6, 0.5, c.get('line1', '').upper(), size=20, font=F('title2'), bold=B('title2'),
                color='FFFFFF', nm='!!l1', wrap=False)
    name = c.get('name', b['name']).upper()
    nsz = fit(name, 'title', 60, 6.6 if not TECH() else 6.8, 30)
    l2 = title_text(s, ML, y + 0.5, 7.0, nsz / 72 * 1.3, name, nsz, color='FFFFFF', nm='!!l2', wrap=False)
    if TECH():
        glow_text(l2, T['light'], 0.12, 35)
    sl = K.text(s, ML, y + 0.6 + nsz / 72 * 1.3, 6.4, 0.5, c.get('slogan', ''), size=14, font=F('body'),
                color='E6EEFF' if TECH() else 'E8F7EE', italic=not TECH())
    lg = logo(s, ctx, x=ML, y=0.45, dark=True) if not TECH() else []
    pres = K.text(s, ML, SH - 0.9, 7.0, 0.4, c.get('presenter', ''), size=11, font=F('semi'), color='C9D6F5'
                  if TECH() else 'D8F2E3')
    K.anims(s, out, 'zoom' if TECH() else 'fly_r', 0.0, 0.2, 1.0)
    if kick is not None:
        K.anim(s, kick, 'fade', 0.5, 0.5)
    K.anims(s, lg, 'fade', 0.3, 0.05, 0.5)
    K.anim(s, l1, 'wipe_l', 0.7, 0.6)
    K.anim(s, l2, 'rise', 0.9, 0.8)
    K.anim(s, sl, 'fade', 1.4, 0.6)
    K.anim(s, pres, 'fade', 1.6, 0.5)
    K.notes(s, c.get('notes', 'Slide bìa.'))
    return s


# ================================================================== 2. letter from the leader
def letter(prs, ctx, c):
    s = _new(prs, ctx)
    bg_light(s, ctx)
    hd = logo(s, ctx)
    panel = K.rect(s, 0, 0, 4.9, SH, nm='!!panel')
    K.grad_fill(panel, [(0, T['primary'], 100), (100, T['dark'], 100)], angle=60)
    if TECH():
        K.picture(s, ctx['art']['ring'], -1.6, 2.9, 5.2, 5.2, alpha=45)
    pp = K.picture(s, pic(ctx, c['portrait'], prefer_cut=True), 0.45, 1.2, 4.1, SH - 1.2, fit='contain', nm='!!hero')
    tag = card(s, 0.55, SH - 1.45, 3.9, 0.95, radius=0.12)
    nm_ = K.text(s, 0.75, SH - 1.4, 3.6, 0.45, c.get('name', ''), size=15, font=F('bold'), bold=B('bold'),
                 color=T['ink'], anchor='m')
    rl = K.text(s, 0.75, SH - 0.98, 3.6, 0.35, c.get('role', ''), size=10.5, font=F('body'), color=T['primary'],
                anchor='m')
    x0, w = 5.6, SW - ML - 5.6
    tt = title_text(s, x0, 1.0, w, 0.7, c.get('title', 'Thư ngỏ').upper(), 34, color=T['primary'] if TECH()
                    else T['ink'], nm='!!ttl', wrap=False)
    qm = K.text(s, x0 - 0.1, 1.55, 1.0, 1.0, '“', size=80, font=F('title'), color=T['pale'])
    body = rich(s, x0, 2.0, w, 3.7, c['text'], size=12.5, line_sp=1.35, space_after=8, align='j')
    sig = K.text(s, SW - ML - 4.0, SH - 1.45, 4.0, 0.7, c.get('signature', ''), size=30, font=F('script'),
                 color=T['ink'], align='r', anchor='m')
    K.anim(s, panel, 'wipe_r', 0.0, 0.6)
    K.anim(s, pp, 'rise', 0.3, 0.8)
    K.anims(s, [tag, nm_, rl], 'fade', 0.8, 0.05, 0.4)
    K.anims(s, hd, 'fade', 0.2, 0.0, 0.4)
    K.anim(s, tt, 'wipe_l', 0.5, 0.6)
    K.anim(s, qm, 'zoom', 0.7, 0.4)
    K.anim(s, body, 'fade', 0.9, 0.8)
    K.anim(s, sig, 'wipe_l', 1.6, 0.8)
    return s


# ================================================================== 3. agenda
def agenda(prs, ctx, c):
    items = c['items']
    if TECH():
        s = _new(prs, ctx, dark=True)
        bg_dark(s, ctx)
        em = K.picture(s, ctx['art']['emblem'], SW / 2 - 1.7, 2.2, 3.4, 3.4, nm='!!emblem')
        tt = title_text(s, 0, 0.6, SW, 0.8, c.get('title', 'Mục lục').upper(), 34, color='FFFFFF', align='c',
                        nm='!!ttl', wrap=False)
        glow_text(tt, T['light'], 0.1, 30)
        n = len(items)
        left = (n + 1) // 2
        out = []
        for i, it in enumerate(items):
            side = 0 if i < left else 1
            k = i if side == 0 else i - left
            rows = left if side == 0 else n - left
            y = 2.15 + k * (3.9 / max(1, rows - 1) if rows > 1 else 0) - 0.1
            x = 0.9 if side == 0 else SW / 2 + 1.95
            w = SW / 2 - 1.95 - 0.9
            num = title_text(s, x if side == 0 else x, y, 1.0, 0.7, f'{i + 1:02d}', 34, color=T['light'],
                             anchor='m', align='l', wrap=False)
            glow_text(num, T['light'], 0.08, 30)
            lb = K.text(s, x + 1.05, y, w - 1.05, 0.7, it.upper(), size=12.5, font=F('title2'), bold=B('title2'),
                        color='FFFFFF', anchor='m', line_sp=1.05)
            ln = K.line(s, x, y + 0.78, x + w, y + 0.78, T['light'], 0.75, alpha=40)
            out.append([num, lb, ln])
        K.anim(s, em, 'turn', 0.0, 1.0)
        K.anim(s, tt, 'zoom', 0.3, 0.6)
        K.anims(s, out, 'fade', 0.6, 0.15, 0.5)
        return s
    s = _new(prs, ctx)
    bg_light(s, ctx)
    hd = header(s, ctx, c.get('title', 'Mục lục'), c.get('sub', ''))
    ph = photo(s, ctx, c['photo'], 8.2, 0.0, 5.6, SH, focus=c.get('focus', (0.5, 0.5)), nm='!!ph1')
    out = []
    n = len(items)
    h = min(0.82, 5.2 / n - 0.12)
    for i, it in enumerate(items):
        y = 1.75 + i * (h + 0.14)
        x = ML + i * 0.12
        cd = K.rect(s, x, y, 7.0, h, fill=T['card'], shp=MSO_SHAPE.PARALLELOGRAM)
        cd.adjustments[0] = 0.18
        K.shadow(cd, blur=0.25, dist=0.05, alpha=14, color='0B1F4D', angle=90)
        K._ensure_ln_before_effects(cd)
        nb = K.rect(s, x + 0.2, y + 0.1, 0.95, h - 0.2, fill=T['primary'], shp=MSO_SHAPE.PARALLELOGRAM)
        nb.adjustments[0] = 0.25
        K.shape_text(nb, f'{i + 1:02d}', size=16, color='FFFFFF', font=F('title'), margin=(0, 0))
        lb = K.text(s, x + 1.35, y, 5.4, h, it, size=14, font=F('semi'), color=T['ink'], anchor='m')
        out.append([cd, nb, lb])
    K.anims(s, hd, 'fade', 0.0, 0.05, 0.4)
    K.anim(s, ph, 'fly_r', 0.1, 0.8)
    K.anims(s, out, 'fly_l', 0.4, 0.12, 0.5)
    return s


# ================================================================== 4. section
def section(prs, ctx, c):
    s = _new(prs, ctx, dark=True, tr=1.3)
    K.picture(s, ctx['art']['sec_' + c['photo']], 0, 0, SW, SH, nm='!!bg')
    num = f'{c["num"]:02d}' if isinstance(c['num'], int) else str(c['num'])
    if TECH():
        nb = title_text(s, 6.4, 0.6, 6.4, 3.6, num, 230, color='FFFFFF', align='r', nm='!!secnum', wrap=False)
        for r in nb._element.iter(K.qn('a:rPr')):
            sf = r.find(K.qn('a:solidFill'))
            if sf is not None:
                r.remove(sf)
            r.insert(0, K.X(f'<a:ln w="25400"><a:solidFill><a:srgbClr val="{T["light"]}"/></a:solidFill></a:ln>'))
            r.insert(1, K.X(f'<a:solidFill><a:srgbClr val="{T["light"]}"><a:alpha val="22000"/></a:srgbClr>'
                            f'</a:solidFill>'))
    else:
        nb = title_text(s, 6.2, 0.9, 6.6, 3.4, num, 220, color='FFFFFF', align='r', nm='!!secnum', wrap=False)
        for r in nb._element.iter(K.qn('a:rPr')):
            sf = r.find(K.qn('a:solidFill'))
            if sf is not None:
                c_ = sf.find(K.qn('a:srgbClr'))
                K.etree.SubElement(c_, K.qn('a:alpha')).set('val', '45000')
    cut = None
    if c.get('cut'):                                   # cut-out person / object standing on the left
        cut = K.picture(s, pic(ctx, c['cut'], prefer_cut=True), 0.0, 0.9, 5.6, SH - 0.9, fit='contain', nm='!!cut')
    t = c['title'].upper() if TECH() else c['title']
    lines = t.split('\n')
    size = 44
    while max(K.text_w(x, F('title'), size) for x in lines) > (6.8 if cut is not None else 7.6) and size > 26:
        size -= 2
    tt = title_text(s, SW - ML - 8.0, 4.1, 8.0, size / 72 * 1.25 * len(lines) + 0.1, t, size, color='FFFFFF',
                    align='r', nm='!!sectitle', line_sp=0.95)
    if TECH():
        glow_text(tt, T['light'], 0.1, 30)
    ds = K.text(s, SW - ML - 7.0, 4.25 + size / 72 * 1.25 * len(lines), 7.0, 0.9, c.get('desc', ''), size=13,
                font=F('body'), color='E6EEFF', align='r', line_sp=1.2) if c.get('desc') else None
    lg = logo(s, ctx, dark=True)
    if cut is not None:
        K.anim(s, cut, 'fly_l', 0.2, 0.8)
    K.anim(s, nb, 'zoom', 0.1, 0.9)
    K.anim(s, tt, 'rise', 0.5, 0.7)
    if ds is not None:
        K.anim(s, ds, 'fade', 0.9, 0.6)
    K.anims(s, lg, 'fade', 0.2, 0.0, 0.4)
    return s


# ================================================================== 5. overview + stats
def overview(prs, ctx, c):
    s = _new(prs, ctx)
    bg_light(s, ctx)
    hd = header(s, ctx, c['title'], c.get('sub', ''), c.get('tag', ''), c.get('num'))
    tx = rich(s, ML, 1.95, 6.3, 2.2, c['text'], size=12.5, line_sp=1.35, space_after=8)
    ph = photo(s, ctx, c['photo'], 7.4, 1.75, SW - ML - 7.4, 3.0, focus=c.get('focus', (0.5, 0.5)), nm='!!ph1')
    stats = c.get('stats', [])[:4]
    n = len(stats)
    gap = 0.25
    w = (SW - 2 * ML - gap * (n - 1)) / max(1, n)
    out = []
    for i, (ic, v, u, lab) in enumerate(stats):
        x = ML + i * (w + gap)
        y = 5.05
        cd = card(s, x, y, w, 1.65)
        bd = icon_badge(s, ic, x + 0.5, y + 0.5, 0.56)
        vt = title_text(s, x + 0.95, y + 0.12, w - 1.0, 0.75, [[(v, {}), (u, {'size': 13})]], 26,
                        color=T['primary'], anchor='m', wrap=False)
        lt = K.text(s, x + 0.25, y + 0.95, w - 0.4, 0.6, lab, size=10.5, font=F('body'), color=T['text'],
                    line_sp=1.1)
        out.append([cd] + bd + [vt, lt])
    K.anims(s, hd, 'fade', 0.0, 0.05, 0.4)
    K.anim(s, tx, 'fade', 0.3, 0.6)
    K.anim(s, ph, 'zoom', 0.4, 0.7)
    K.anims(s, out, 'rise', 0.8, 0.15, 0.5)
    return s


# ================================================================== 6. milestones / steps
def milestones(prs, ctx, c):
    s = _new(prs, ctx)
    bg_light(s, ctx)
    hd = header(s, ctx, c['title'], c.get('sub', ''), c.get('tag', ''), c.get('num'))
    ev = c['events']
    n = len(ev)
    x0, x1 = ML + 0.3, SW - ML - 0.3
    step = (x1 - x0) / n
    y = 4.2                       # rail; even items: year above / card below, odd items: card above / year below
    ch = 1.65
    rail = K.rect(s, x0, y - 0.03, x1 - x0, 0.06)
    K.grad_fill(rail, [(0, T['light'], 100), (100, T['primary'], 100)], angle=0)
    out = []
    for i, (yr, tt, tx) in enumerate(ev):
        cx = x0 + step * (i + 0.5)
        below = i % 2 == 0
        cw = step - 0.15
        cy = y + 0.55 if below else y - 0.55 - ch
        stem = K.line(s, cx, y + (0.17 if below else -0.17), cx, cy if below else cy + ch, T['primary'], 1.25,
                      dash='dash')
        dot = K.oval(s, cx - 0.17, y - 0.17, 0.34, fill='FFFFFF', line=T['primary'], lw=3)
        yt = title_text(s, cx - step / 2, (y - 0.85) if below else (y + 0.25), step, 0.6, yr,
                        fit(yr, 'title', 22, step - 0.2, 12), color=T['primary'], align='c', anchor='m', wrap=False)
        cd = card(s, cx - cw / 2, cy, cw, ch)
        t1 = K.text(s, cx - cw / 2 + 0.15, cy + 0.14, cw - 0.3, 0.45, tt, size=12.5, font=F('bold'),
                    bold=B('bold'), color=T['ink'], align='c', anchor='m', line_sp=1.05)
        t2 = K.text(s, cx - cw / 2 + 0.15, cy + 0.62, cw - 0.3, ch - 0.7, tx, size=10.5, font=F('body'),
                    color=T['text'], align='c', line_sp=1.15)
        out.append([dot, yt, stem, cd, t1, t2])
    K.anims(s, hd, 'fade', 0.0, 0.05, 0.4)
    K.anim(s, rail, 'wipe_l', 0.3, 0.9)
    K.anims(s, out, 'rise', 0.5, 0.18, 0.45)
    return s


# ================================================================== 7. org chart
def org_chart(prs, ctx, c):
    s = _new(prs, ctx)
    bg_light(s, ctx)
    hd = header(s, ctx, c['title'], c.get('sub', ''), c.get('tag', ''), c.get('num'))
    levels = c['levels']           # [[top], [second], [third...]] — each level centred
    out = []
    ys = [1.95, 2.95, 3.95, 5.35]
    prev = None
    for li, names in enumerate(levels[:4]):
        n = len(names)
        w = min(3.3, (SW - 2 * ML - 0.2 * (n - 1)) / n)
        tot = n * w + (n - 1) * 0.2
        x0 = (SW - tot) / 2
        y = ys[li]
        h = 0.62 if li < 3 else 0.95
        centers = []
        boxes = []
        for i, nm_ in enumerate(names):
            x = x0 + i * (w + 0.2)
            fill = T['dark'] if li == 0 else (T['primary'] if li == 1 else (T['light'] if li == 2 else T['card']))
            col = 'FFFFFF' if li < 2 else (T['dark'] if li == 2 else T['ink'])
            bx = K.rect(s, x, y, w, h, fill=fill, radius=0.1, line=T['line'] if li == 3 else None)
            K.shadow(bx, blur=0.2, dist=0.04, alpha=14, color='0B1F4D', angle=90)
            K._ensure_ln_before_effects(bx)
            K.shape_text(bx, nm_, size=11 if li < 3 else 10, color=col, font=F('semi'), margin=(0.06, 0.03),
                         line_sp=1.05)
            boxes.append(bx)
            centers.append(x + w / 2)
        lines_ = []
        if prev:
            py, pcs = prev
            midy = (py + y) / 2
            pc = sum(pcs) / len(pcs)
            lines_.append(K.line(s, pc, py, pc, midy, T['primary'], 1.25))
            lines_.append(K.line(s, min(centers), midy, max(centers), midy, T['primary'], 1.25))
            for cx in centers:
                lines_.append(K.line(s, cx, midy, cx, y, T['primary'], 1.25))
        out.append((lines_, boxes))
        prev = (y + h, centers)
    K.anims(s, hd, 'fade', 0.0, 0.05, 0.4)
    t = 0.3
    for lines_, boxes in out:
        K.anims(s, lines_, 'wipe_d', t, 0.03, 0.3)
        K.anims(s, boxes, 'zoom', t + 0.2, 0.08, 0.4)
        t += 0.5
    return s


# ================================================================== 8. pillars (3 cards)
def pillars(prs, ctx, c):
    s = _new(prs, ctx)
    bg_light(s, ctx)
    hd = header(s, ctx, c['title'], c.get('sub', ''), c.get('tag', ''), c.get('num'))
    lead = rich(s, ML, 1.85, SW - 2 * ML - (4.2 if c.get('photo') else 0), 0.8, c.get('lead', ''), size=12.5) \
        if c.get('lead') else None
    ph = photo(s, ctx, c['photo'], SW - ML - 3.9, 1.0 if not TECH() else 0.95, 3.9, 1.65,
               focus=c.get('focus', (0.5, 0.5))) if c.get('photo') else None
    items = c['items'][:3]
    n = len(items)
    gap = 0.3
    w = (SW - 2 * ML - gap * (n - 1)) / n
    est = max(sum(K.n_lines(plain(t), F('body'), 12, w - 0.8) * 12 * 1.25 / 72 + 0.11 for t in it[2])
              for it in items)
    chh = min(3.95, max(3.2, 1.3 + est + 0.4))          # card height follows the longest bullet list
    y = 2.95
    out = []
    for i, (ic, tt, lines) in enumerate(items):
        x = ML + i * (w + gap)
        cd = card(s, x, y, w, chh)
        top = K.rect(s, x, y, w, 0.12, fill=T['primary'] if i != 1 else T['light'])
        bd = icon_badge(s, ic, x + 0.6, y + 0.65, 0.66)
        t1 = K.text(s, x + 1.05, y + 0.32, w - 1.2, 0.7, tt, size=14.5, font=F('bold'), bold=B('bold'), color=T['ink'],
                    anchor='m', line_sp=1.05)
        bl = rich(s, x + 0.3, y + 1.3, w - 0.55, chh - 1.4, '\n'.join(lines), size=12, line_sp=1.25, space_after=8,
                  bullets=True)
        out.append([cd, top] + bd + [t1, bl])
    K.anims(s, hd, 'fade', 0.0, 0.05, 0.4)
    if lead is not None:
        K.anim(s, lead, 'fade', 0.3, 0.5)
    if ph is not None:
        K.anim(s, ph, 'fly_r', 0.3, 0.6)
    K.anims(s, out, 'rise', 0.5, 0.2, 0.5)
    return s


# ================================================================== 9. services list + photo
def services(prs, ctx, c):
    s = _new(prs, ctx, dark=bool(c.get('dark')))
    if c.get('dark'):
        bg_dark(s, ctx, 'dark_bg' if TECH() else 'cover_bg')
    else:
        bg_light(s, ctx)
    hd = header(s, ctx, c['title'], c.get('sub', ''), c.get('tag', ''), c.get('num'), dark=bool(c.get('dark')))
    ph = photo(s, ctx, c['photo'], 8.0, 1.75, SW - ML - 8.0, 5.0, focus=c.get('focus', (0.5, 0.5)), nm='!!ph1')
    items = c['items'][:5]
    n = len(items)
    h = min(0.95, (5.0 - 0.15 * (n - 1)) / n)
    out = []
    for i, (ic, tt, tx) in enumerate(items):
        y = 1.75 + i * (h + 0.15)
        cd = card(s, ML, y, 7.1, h, radius=0.12, fill=T['card'] if not c.get('dark') else 'FFFFFF', alpha=None)
        acc = K.rect(s, ML, y, 0.1, h, fill=T['primary'] if i % 2 == 0 else T['light'])
        bd = icon_badge(s, ic, ML + 0.6, y + h / 2, 0.52)
        t1 = K.text(s, ML + 1.05, y + 0.06, 5.9, 0.38, tt, size=13, font=F('bold'), bold=B('bold'), color=T['ink'],
                    anchor='m')
        t2 = K.text(s, ML + 1.05, y + 0.42, 5.9, h - 0.45, tx, size=10.5, font=F('body'), color=T['text'],
                    line_sp=1.1)
        out.append([cd, acc] + bd + [t1, t2])
    K.anims(s, hd, 'fade', 0.0, 0.05, 0.4)
    K.anim(s, ph, 'fly_r', 0.2, 0.7)
    K.anims(s, out, 'fly_l', 0.4, 0.13, 0.45)
    return s


# ================================================================== 10. project / package cards
def projects(prs, ctx, c):
    s = _new(prs, ctx)
    bg_light(s, ctx)
    hd = header(s, ctx, c['title'], c.get('sub', ''), c.get('tag', ''), c.get('num'))
    items = c['items'][:3]
    n = len(items)
    gap = 0.3
    w = (SW - 2 * ML - gap * (n - 1)) / n
    out = []
    for i, it in enumerate(items):
        x = ML + i * (w + gap)
        y = 1.85
        cd = card(s, x, y, w, 4.6)
        ph = K.picture(s, pic(ctx, it['photo']), x, y, w, 2.45, focus=tuple(it.get('focus', (0.5, 0.5))),
                       geom='round2SameRect')
        g = ph._element.spPr.find(K.qn('a:prstGeom'))
        av = g.find(K.qn('a:avLst'))
        for nm_, v in (('adj1', int(0.14 / 2.45 * 100000)), ('adj2', 0)):
            gd = K.etree.SubElement(av, K.qn('a:gd'))
            gd.set('name', nm_)
            gd.set('fmla', f'val {v}')
        if it.get('badge'):
            bw = K.text_w(it['badge'], F('semi'), 9.5) + 0.4
            bg_ = K.rect(s, x + 0.2, y + 0.2, bw, 0.32, fill=T['primary'], radius=0.16)
            K.shape_text(bg_, it['badge'], size=9.5, color='FFFFFF', font=F('semi'), margin=(0.06, 0.0))
        else:
            bg_ = None
        t1 = title_text(s, x + 0.25, y + 2.6, w - 0.5, 0.55, it['name'].upper(), fit(it['name'].upper(), 'title', 16,
                                                                                     w - 0.5, 11),
                        color=T['primary'] if TECH() else T['ink'], anchor='m', wrap=False)
        meta = K.text(s, x + 0.25, y + 3.15, w - 0.5, 0.35, it.get('meta', ''), size=10, font=F('semi'),
                      color=T['muted'])
        t2 = rich(s, x + 0.25, y + 3.5, w - 0.5, 1.0, it.get('text', ''), size=11.5, line_sp=1.2)
        out.append([cd, ph] + ([bg_] if bg_ is not None else []) + [t1, meta, t2])
    K.anims(s, hd, 'fade', 0.0, 0.05, 0.4)
    K.anims(s, out, 'rise', 0.3, 0.2, 0.55)
    return s


# ================================================================== 11. big numbers
def big_numbers(prs, ctx, c):
    s = _new(prs, ctx)
    bg_light(s, ctx)
    hd = header(s, ctx, c['title'], c.get('sub', ''), c.get('tag', ''), c.get('num'))
    items = c['items'][:4]
    ph = photo(s, ctx, c['photo'], SW - ML - 4.6, 1.85, 4.6, 4.9, focus=c.get('focus', (0.5, 0.5)), nm='!!ph1') \
        if c.get('photo') else None
    wtot = (SW - 2 * ML - (4.9 if ph is not None else 0))
    cw = (wtot - 0.3) / 2
    out = []
    for i, (ic, v, u, lab) in enumerate(items):
        x = ML + (i % 2) * (cw + 0.3)
        y = 1.85 + (i // 2) * 2.55
        fill = T['primary'] if i in (0, 3) else T['card']
        col = 'FFFFFF' if i in (0, 3) else T['ink']
        cd = card(s, x, y, cw, 2.3, fill=fill)
        bd = K.icon(s, ic, x + 0.3, y + 0.3, 0.45, 'FFFFFF' if i in (0, 3) else T['primary'])
        vt = title_text(s, x + 0.3, y + 0.8, cw - 0.5, 0.85, [[(v, {}), (u, {'size': 18})]], 40, color=col,
                        anchor='m', wrap=False)
        lt = K.text(s, x + 0.3, y + 1.65, cw - 0.5, 0.55, lab, size=11, font=F('body'), color=col, line_sp=1.1)
        out.append([cd, bd, vt, lt])
    K.anims(s, hd, 'fade', 0.0, 0.05, 0.4)
    if ph is not None:
        K.anim(s, ph, 'fly_r', 0.2, 0.7)
    K.anims(s, out, 'zoom', 0.4, 0.15, 0.5)
    return s


# ================================================================== 12. app / digital platform (phone mock-up in shapes)
def phone_app(prs, ctx, c):
    s = _new(prs, ctx)
    bg_light(s, ctx)
    hd = header(s, ctx, c['title'], c.get('sub', ''), c.get('tag', ''), c.get('num'))
    px, py, pw, ph_ = 8.6, 1.35, 2.85, 5.75
    acc = K.rect(s, px - 1.1, py + 0.6, 1.4, 4.6, fill=T['accent'], shp=MSO_SHAPE.PARALLELOGRAM)
    acc.adjustments[0] = 0.35
    body = K.rect(s, px, py, pw, ph_, fill='1A1F2B', radius=0.38)
    K.shadow(body, blur=0.45, dist=0.15, alpha=35, color='000000', angle=90)
    K._ensure_ln_before_effects(body)
    scr = K.rect(s, px + 0.1, py + 0.1, pw - 0.2, ph_ - 0.2, fill='FFFFFF', radius=0.3)
    notch = K.rect(s, px + pw / 2 - 0.45, py + 0.14, 0.9, 0.16, fill='1A1F2B', radius=0.08)
    top = K.rect(s, px + 0.1, py + 0.1, pw - 0.2, 1.45, fill=T['primary'], shp=MSO_SHAPE.ROUND_2_SAME_RECTANGLE)
    top.adjustments[0] = 0.3 / 1.45
    hi = K.text(s, px + 0.3, py + 0.45, pw - 0.6, 0.3, c.get('app_greeting', 'Xin chào,'), size=9, font=F('body'),
                color='FFFFFF')
    hn = K.text(s, px + 0.3, py + 0.7, pw - 0.6, 0.35, c.get('app_user', 'Nguyễn Văn A'), size=12, font=F('bold'),
                bold=B('bold'), color='FFFFFF')
    bal = K.rect(s, px + 0.25, py + 1.2, pw - 0.5, 0.95, fill='FFFFFF', radius=0.12)
    K.shadow(bal, blur=0.2, dist=0.04, alpha=18, color='000000', angle=90)
    K._ensure_ln_before_effects(bal)
    bt = K.text(s, px + 0.4, py + 1.27, pw - 0.8, 0.28, c.get('app_card_label', 'Hợp đồng đang hiệu lực'), size=8,
                font=F('body'), color=T['muted'])
    bv = K.text(s, px + 0.4, py + 1.52, pw - 0.8, 0.5, c.get('app_card_value', '02'), size=20, font=F('title'),
                bold=B('title'), color=T['primary'], anchor='m')
    tiles = []
    for i, (ic, lb) in enumerate(c.get('app_tiles', [])[:6]):
        tx = px + 0.3 + (i % 3) * ((pw - 0.6) / 3)
        ty = py + 2.4 + (i // 3) * 0.95
        tiles += K.icon_badge(s, ic, tx + (pw - 0.6) / 6, ty + 0.25, 0.42, T['pale'], T['primary'], icon_scale=0.55)
        tiles.append(K.text(s, tx, ty + 0.5, (pw - 0.6) / 3, 0.3, lb, size=7, font=F('body'), color=T['ink'],
                            align='c'))
    rows = []
    for i in range(2):
        ry = py + 4.45 + i * 0.5
        rows.append(K.rect(s, px + 0.3, ry, pw - 0.6, 0.38, fill=T['pale'], radius=0.08))
    phone = [acc, body, scr, notch, top, hi, hn, bal, bt, bv] + tiles + rows
    items = c['items'][:6]
    out = []
    for i, t in enumerate(items):
        y = 1.95 + i * 0.78
        bd = K.icon(s, 'circle-check', ML, y + 0.17, 0.36, T['primary'])
        tb = rich(s, ML + 0.55, y, 6.9, 0.7, t, size=12.5, line_sp=1.15, space_after=0, anchor='m')
        out.append([bd, tb])
    K.anims(s, hd, 'fade', 0.0, 0.05, 0.4)
    K.anims(s, phone, 'fade', 0.2, 0.02, 0.4)
    K.anims(s, out, 'fly_l', 0.6, 0.12, 0.45)
    return s


# ================================================================== 13. chart + KPIs
def chart(prs, ctx, c):
    s = _new(prs, ctx)
    bg_light(s, ctx)
    hd = header(s, ctx, c['title'], c.get('sub', ''), c.get('tag', ''), c.get('num'))
    cd = card(s, ML, 1.85, 8.0, 4.95)
    ch = c['chart']
    if ch.get('kind') == 'line':
        gf = K.chart_line(s, ML + 0.25, 2.05, 7.5, 4.3, ch['cats'], [(n, v) for n, v in ch['series']],
                          colors=T['chart'], legend=True, num_fmt=ch.get('fmt', '0'), markers=True, font_size=10.5,
                          legend_pos='b')
    else:
        gf = K.chart_bar(s, ML + 0.25, 2.05, 7.5, 4.3, ch['cats'], [(n, v) for n, v in ch['series']],
                         colors=T['chart'], gap=60, overlap=-10, labels=True, num_fmt=ch.get('fmt', '0'), legend=True,
                         font_size=10.5, label_color=T['ink'], max_v=ch.get('max'))
    g = gf.chart
    g.font.name = F('body')
    note = K.text(s, ML + 0.3, 6.35, 7.4, 0.35, c.get('note', ''), size=9.5, font=F('body'), italic=True,
                  color=T['muted'])
    kp = []
    x0 = 8.9
    w = SW - ML - x0
    for i, (v, u, t) in enumerate(c.get('kpis', [])[:3]):
        y = 1.85 + i * 1.7
        b = card(s, x0, y, w, 1.5, fill=T['primary'] if i == 0 else T['card'])
        col = 'FFFFFF' if i == 0 else T['primary']
        vt = title_text(s, x0 + 0.3, y + 0.12, w - 0.5, 0.75, [[(v, {}), (u, {'size': 15})]], 30, color=col,
                        anchor='m', wrap=False)
        lt = K.text(s, x0 + 0.3, y + 0.85, w - 0.5, 0.55, t, size=10.5, font=F('body'),
                    color='FFFFFF' if i == 0 else T['text'], line_sp=1.1)
        kp.append([b, vt, lt])
    K.anims(s, hd, 'fade', 0.0, 0.05, 0.4)
    K.anim(s, cd, 'fade', 0.2, 0.4)
    K.anim(s, gf, 'wipe_u' if ch.get('kind') != 'line' else 'wipe_l', 0.4, 1.0)
    K.anim(s, note, 'fade', 1.2, 0.4)
    K.anims(s, kp, 'fly_r', 0.5, 0.2, 0.5)
    return s


# ================================================================== 14. table
def table_slide(prs, ctx, c):
    s = _new(prs, ctx)
    bg_light(s, ctx)
    hd = header(s, ctx, c['title'], c.get('sub', ''), c.get('tag', ''), c.get('num'))
    rows = c['rows']
    w = SW - 2 * ML
    cw = c.get('col_w')
    if cw:
        k = w / sum(cw)
        cw = [v * k for v in cw]
    rh = min(0.72, 4.6 / len(rows))
    gf = K.table(s, ML, 1.95, w, rh * len(rows), rows, col_w=cw, header_fill=T['primary'], header_color='FFFFFF',
                 body_fill=T['card'], body_fill2=T['pale'], body_color=T['ink'], line_color=T['line'], size=11.5,
                 header_size=12, row_h=rh)
    tbl = gf.table
    for r, row in enumerate(tbl.rows):
        for ci, cell in enumerate(row.cells):
            for p in cell.text_frame.paragraphs:
                for run in p.runs:
                    rPr = run._r.get_or_add_rPr()
                    face = F('semi') if (r == 0 or ci == 0) else F('body')
                    for tag in ('a:latin', 'a:ea', 'a:cs'):
                        el = rPr.find(K.qn(tag))
                        if el is not None:
                            el.set('typeface', face)
                    run.font.bold = False
    note = K.text(s, ML, 1.95 + rh * len(rows) + 0.15, w, 0.35, c.get('note', ''), size=9.5, font=F('body'),
                  italic=True, color=T['muted'])
    K.anims(s, hd, 'fade', 0.0, 0.05, 0.4)
    K.anim(s, gf, 'wipe_d', 0.3, 0.9)
    K.anim(s, note, 'fade', 1.1, 0.4)
    return s


# ================================================================== 15. thanks
def thanks(prs, ctx, c):
    s = _new(prs, ctx, dark=True, tr=1.3)
    out = []
    if TECH():
        bg_dark(s, ctx, 'cover_bg')
        ed = 4.2 if c.get('qr') else 4.6
        out.append(K.picture(s, ctx['art']['emblem'], 10.5 - ed / 2, 0.75 if c.get('qr') else 1.1, ed, ed,
                             nm='!!emblem'))
    else:
        bg_dark(s, ctx, 'cover_bg')
        out.append(photo(s, ctx, c['photo'], 7.2, 0.0, 6.2, SH, focus=c.get('focus', (0.5, 0.5)), nm='!!ph1'))
    t = c.get('title', 'Xin cảm ơn').upper()
    size = fit(t, 'title', 64, 6.8, 30)
    tt = title_text(s, ML, 1.6, 7.0, size / 72 * 1.3, t, size, color='FFFFFF', nm='!!l2', wrap=False)
    if TECH():
        glow_text(tt, T['light'], 0.12, 35)
    sub = K.text(s, ML, 1.6 + size / 72 * 1.3, 6.6, 0.5, c.get('sub', ''), size=16, font=F('body'),
                 color='E6EEFF' if TECH() else 'E8F7EE', italic=not TECH())
    rows = []
    y = 3.6
    for ic, tx in c.get('lines', [])[:5]:
        rows.append([K.icon(s, ic, ML, y + 0.04, 0.32, 'FFFFFF'),
                     K.text(s, ML + 0.5, y, 6.2, 0.42, tx, size=13, font=F('semi'), color='FFFFFF', anchor='m')])
        y += 0.52
    lg = logo(s, ctx, x=ML, y=0.45, dark=True)
    qr = []
    if c.get('qr'):
        qx, qy = (9.65, 5.0) if TECH() else (ML + 4.9, 5.0)
        qc = K.rect(s, qx, qy, 1.7, 2.0, fill='FFFFFF', radius=0.12)
        K.shadow(qc, blur=0.35, dist=0.1, alpha=35, color='000000', angle=90)
        K._ensure_ln_before_effects(qc)
        qr = [qc, K.picture(s, pic(ctx, c['qr']), qx + 0.15, qy + 0.13, 1.4, 1.4),
              K.text(s, qx, qy + 1.55, 1.7, 0.35, c.get('qr_label', 'Quét mã QR'), size=9.5, font=F('semi'),
                     color=T['ink'], align='c', anchor='m')]
    K.anims(s, out, 'zoom' if TECH() else 'fly_r', 0.0, 0.0, 1.0)
    K.anim(s, tt, 'rise', 0.4, 0.8)
    if qr:
        K.anims(s, [qr], 'zoom', 1.5, 0.0, 0.5)
    K.anim(s, sub, 'fade', 0.9, 0.5)
    K.anims(s, rows, 'wipe_l', 1.1, 0.12, 0.4)
    K.anims(s, lg, 'fade', 0.2, 0.0, 0.4)
    return s


# ================================================================== helpers for image-led slides
def _img_size(path):
    from PIL import Image
    return Image.open(path).size


def _contain(s, path, x, y, w, h, nm=None, valign='m'):
    """Picture scaled to fit inside the box; returns (shape, (x, y, w, h) actually used)."""
    iw, ih = _img_size(path)
    r = min(w / iw, h / ih)
    pw, ph = iw * r, ih * r
    px = x + (w - pw) / 2
    py = y + {'t': 0, 'm': (h - ph) / 2, 'b': h - ph}[valign]
    return K.picture(s, path, px, py, pw, ph, nm=nm), (px, py, pw, ph)


def browser(s, path, x, y, w, url='', nm=None):
    """Screenshot inside a browser window (title bar with three dots and the address); height follows the image."""
    iw, ih = _img_size(path)
    bar = 0.34
    h = (w - 0.1) * ih / iw + bar + 0.05
    fr = K.rect(s, x, y, w, h, fill='FFFFFF', radius=0.1, line=T['line'], lw=0.75, nm=nm)
    K.shadow(fr, blur=0.45, dist=0.12, alpha=28, color='0B1F2A', angle=90)
    K._ensure_ln_before_effects(fr)
    out = [fr] + [K.oval(s, x + 0.16 + k * 0.17, y + 0.12, 0.1, fill=col)
                  for k, col in enumerate(('FF5F57', 'FEBC2E', '28C840'))]
    if url:
        ub = K.rect(s, x + 0.75, y + 0.075, min(3.4, w - 1.0), 0.2, fill=T['bg'], radius=0.1)
        K.shape_text(ub, url, size=7.5, font=F('body'), color=T['muted'], align='l', margin=(0.12, 0))
        out.append(ub)
    out.append(K.picture(s, path, x + 0.05, y + bar, w - 0.1, h - bar - 0.05))
    return out, h


def _callout(s, x, y, w, h, ic, title, text, compact=False):
    cd = card(s, x, y, w, h)
    if compact:          # icon left, title + text to the right (stacked lists)
        bd = icon_badge(s, ic, x + 0.5, y + h / 2, 0.58)
        t1 = K.text(s, x + 0.95, y + 0.1, w - 1.1, 0.4, title, size=13, font=F('bold'), bold=B('bold'),
                    color=T['ink'], anchor='m')
        t2 = rich(s, x + 0.95, y + 0.5, w - 1.1, h - 0.55, text, size=10.5, line_sp=1.12, space_after=0)
    else:                # icon + title on top, text below (2 + 2 around an image)
        bd = icon_badge(s, ic, x + 0.55, y + 0.55, 0.62)
        t1 = K.text(s, x + 1.0, y + 0.18, w - 1.15, 0.75, title, size=14, font=F('bold'), bold=B('bold'),
                    color=T['ink'], anchor='m', line_sp=1.0)
        t2 = rich(s, x + 0.28, y + 1.05, w - 0.5, h - 1.12, text, size=11, line_sp=1.18, space_after=0)
    return [cd] + bd + [t1, t2]


# ================================================================== 16. hub: emblem + 4 pillars
def hub(prs, ctx, c):
    """Brand emblem in the centre, four pillar cards in the corners linked to it (ecosystem / 4 business lines)."""
    import math
    s = _new(prs, ctx)
    bg_light(s, ctx)
    hd = header(s, ctx, c['title'], c.get('sub', ''), c.get('tag', ''), c.get('num'))
    cx, cy, d = SW / 2, 4.42, 2.75
    rr = d * 0.6
    glow = K.glow_img(s, cx, cy, d * 1.7, T['light'], 40)
    ring = K.oval(s, cx - rr, cy - rr, 2 * rr, fill=None, line=T['primary'], lw=1.25, dash='dash')
    em = K.picture(s, ctx['art']['emblem'], cx - d / 2, cy - d / 2, d, d, nm='!!emblem')
    cw, ch = 4.1, 2.2
    pos = [(ML, 1.95), (SW - ML - cw, 1.95), (ML, 4.62), (SW - ML - cw, 4.62)]
    links, cards_ = [], []
    for i, (ic, tt, tx) in enumerate(c['items'][:4]):
        x, y = pos[i]
        ex, ey = (x + cw) if i % 2 == 0 else x, y + ch / 2
        ang = math.atan2(ey - cy, ex - cx)
        rx, ry = cx + math.cos(ang) * rr, cy + math.sin(ang) * rr
        links.append([K.line(s, ex, ey, rx, ry, T['primary'], 1.25),
                      K.oval(s, rx - 0.08, ry - 0.08, 0.16, fill=T['light'], line='FFFFFF', lw=1.25)])
        cd = card(s, x, y, cw, ch)
        bd = icon_badge(s, ic, x + 0.6, y + 0.6, 0.66)
        nb = title_text(s, x + cw - 1.05, y + 0.12, 0.85, 0.55, f'{i + 1:02d}', 24, color=K.mix(T['light'], 'FFFFFF',
                                                                                              0.45),
                        align='r', anchor='m', wrap=False)
        t1 = K.text(s, x + 1.1, y + 0.22, cw - 2.0, 0.75, tt, size=15, font=F('bold'), bold=B('bold'),
                    color=T['ink'], anchor='m', line_sp=1.0)
        t2 = rich(s, x + 0.3, y + 1.12, cw - 0.55, ch - 1.2, tx, size=11.5, line_sp=1.2, space_after=0)
        cards_.append([cd] + bd + [nb, t1, t2])
    K.anims(s, hd, 'fade', 0.0, 0.05, 0.4)
    K.anim(s, glow, 'fade', 0.1, 0.6)
    K.anim(s, em, 'turn', 0.2, 0.9)
    K.anim(s, ring, 'zoom', 0.6, 0.5)
    K.anims(s, links, 'fade', 0.9, 0.1, 0.4)
    K.anims(s, cards_, 'zoom', 1.0, 0.18, 0.5)
    return s


# ================================================================== 17. spotlight: hero image + callouts
def spotlight(prs, ctx, c):
    """A hero image (cut-out, illustration or a screenshot in a browser frame) with up to 4 callout cards.
    mode 'center': image in the middle, 2 callouts each side; 'left' / 'right': image on that side, callouts
    stacked on the other. Optional image2 (smaller, overlapping the corner), soft=True fades the image edges."""
    s = _new(prs, ctx)
    bg_light(s, ctx)
    hd = header(s, ctx, c['title'], c.get('sub', ''), c.get('tag', ''), c.get('num'))
    mode = c.get('mode', 'center')
    items = c['items'][:4]
    top, bot = 1.95, SH - 0.45
    if mode == 'center':
        cw = 3.45
        bx, bw = ML + cw + 0.3, SW - 2 * (ML + cw + 0.3)
    else:
        bw = c.get('image_w', 6.55)
        bx = ML if mode == 'left' else SW - ML - bw
    by, bh = top, bot - top
    path = pic(ctx, c['image'], prefer_cut=True)
    glow = K.glow_img(s, bx + bw / 2, by + bh / 2, min(bw, bh) * 1.3, T['light'], 38)
    hero = []
    if c.get('frame') == 'browser':
        fw = bw - (0.9 if c.get('image2') else 0.1)
        parts, fh = browser(s, path, bx + 0.05, 0, fw, c.get('url', ''), nm='!!hero')
        dy = by + max(0.1, (bh - fh) / 2 - (0.35 if c.get('image2') else 0))
        for sh in parts:
            sh.top = sh.top + K.E(dy)
        hero += parts
        box = (bx + 0.05, dy, fw, fh)
    else:
        p, box = _contain(s, path, bx, by, bw, bh, nm='!!hero')
        if c.get('soft'):
            K.soft_edge(p, 0.3)
        hero.append(p)
    if c.get('image2'):
        p2path = pic(ctx, c['image2'], prefer_cut=True)
        w2 = bw * c.get('image2_w', 0.42)
        if c.get('frame2') == 'browser':
            parts2, h2 = browser(s, p2path, 0, 0, w2, c.get('url2', ''))
            x2, y2 = bx + bw - w2, bot - h2
            for sh in parts2:
                sh.left = sh.left + K.E(x2)
                sh.top = sh.top + K.E(y2)
            hero.append(parts2)
        else:
            iw, ih = _img_size(p2path)
            h2 = min(w2 * ih / iw, bh * 0.5)
            w2 = h2 * iw / ih
            p2 = K.picture(s, p2path, bx + bw - w2, bot - h2, w2, h2)
            K.shadow(p2, blur=0.3, dist=0.08, alpha=30, color='0B1F2A', angle=90)
            K._ensure_ln_before_effects(p2)
            hero.append(p2)
    out, links = [], []
    if mode == 'center':
        ch = 2.3
        ys = (2.0, SH - 0.5 - ch)
        for i, (ic, tt, tx) in enumerate(items):
            right = i >= 2
            x = SW - ML - cw if right else ML
            y = ys[i % 2]
            out.append(_callout(s, x, y, cw, ch, ic, tt, tx))
            ex = x if right else x + cw
            ix = box[0] + 0.1 if right is False else box[0] + box[2] - 0.1
            ix = min(max(ix, bx), bx + bw)
            ey = y + ch / 2
            iy = min(max(ey, box[1] + 0.3), box[1] + box[3] - 0.3)
            links.append([K.line(s, ex, ey, ix, iy, T['primary'], 1.0, dash='dash'),
                          K.oval(s, ix - 0.07, iy - 0.07, 0.14, fill=T['light'], line='FFFFFF', lw=1.25)])
    else:
        gx = (bx + bw + 0.45) if mode == 'left' else ML
        gw = (SW - ML - gx) if mode == 'left' else (bx - 0.45 - ML)
        n = len(items)
        gap = 0.17
        ch = (bh - gap * (n - 1)) / n
        for i, (ic, tt, tx) in enumerate(items):
            out.append(_callout(s, gx, by + i * (ch + gap), gw, ch, ic, tt, tx, compact=True))
    K.anims(s, hd, 'fade', 0.0, 0.05, 0.4)
    K.anim(s, glow, 'fade', 0.1, 0.6)
    K.anims(s, hero, 'zoom' if mode == 'center' else ('fly_l' if mode == 'left' else 'fly_r'), 0.2, 0.25, 0.7)
    K.anims(s, links, 'fade', 0.8, 0.1, 0.35)
    K.anims(s, out, 'zoom' if mode == 'center' else ('fly_r' if mode == 'left' else 'fly_l'), 0.7, 0.15, 0.45)
    return s


# ================================================================== 18. cards: photo grid
def cards(prs, ctx, c):
    """Grid of photo cards. Items {photo, title?, text?, badge?, caption?, focus?}; cols (default up to 4).
    Items with only a caption become full-bleed photos with the caption on a dark gradient."""
    import math
    s = _new(prs, ctx)
    bg_light(s, ctx)
    hd = header(s, ctx, c['title'], c.get('sub', ''), c.get('tag', ''), c.get('num'))
    items = c['items']
    n = len(items)
    cols = c.get('cols', min(4, n))
    rows = math.ceil(n / cols)
    gap = 0.25
    top, bot = 1.95, SH - 0.45
    w = (SW - 2 * ML - gap * (cols - 1)) / cols
    h = (bot - top - gap * (rows - 1)) / rows
    out = []
    for i, it in enumerate(items):
        x = ML + (i % cols) * (w + gap)
        y = top + (i // cols) * (h + gap)
        foc = tuple(it.get('focus', (0.5, 0.5)))
        if not (it.get('title') or it.get('text')):
            p = K.picture(s, pic(ctx, it['photo']), x, y, w, h, radius=0.14, focus=foc)
            K.shadow(p, blur=0.3, dist=0.08, alpha=22, color='0B1F2A', angle=90)
            K._ensure_ln_before_effects(p)
            gh = min(h * 0.55, 1.2)
            sh = K.rect(s, x, y + h - gh, w, gh, shp=MSO_SHAPE.ROUND_2_SAME_RECTANGLE, rot=180)
            sh.adjustments[0] = min(0.5, 0.14 / gh)
            K.grad_fill(sh, [(0, T['dark'], 88), (100, T['dark'], 0)], angle=90)
            cap = K.text(s, x + 0.2, y + h - 0.62, w - 0.4, 0.5, it.get('caption', ''), size=it.get('size', 13),
                         font=F('bold'), bold=B('bold'), color='FFFFFF', anchor='m', line_sp=1.0)
            out.append([p, sh, cap])
            continue
        cd = card(s, x, y, w, h)
        ph_h = h * c.get('photo_ratio', 0.55)
        p = K.picture(s, pic(ctx, it['photo']), x, y, w, ph_h, focus=foc, geom='round2SameRect')
        av = p._element.spPr.find(K.qn('a:prstGeom')).find(K.qn('a:avLst'))
        for nm_, v in (('adj1', int(min(0.5, 0.14 / ph_h) * 100000)), ('adj2', 0)):
            gd = K.etree.SubElement(av, K.qn('a:gd'))
            gd.set('name', nm_)
            gd.set('fmla', f'val {v}')
        grp = [cd, p]
        if it.get('badge'):
            bw_ = K.text_w(it['badge'], F('semi'), 9) + 0.36
            b_ = K.rect(s, x + 0.15, y + 0.15, bw_, 0.3, fill=T['primary'], radius=0.15)
            K.shape_text(b_, it['badge'], size=9, color='FFFFFF', font=F('semi'), margin=(0.05, 0.0))
            grp.append(b_)
        ty = y + ph_h + 0.12
        grp.append(K.text(s, x + 0.2, ty, w - 0.4, 0.42, it.get('title', ''), size=c.get('title_size', 13),
                          font=F('bold'), bold=B('bold'), color=T['ink'], anchor='m', line_sp=1.0))
        if it.get('text'):
            grp.append(rich(s, x + 0.2, ty + 0.45, w - 0.4, h - ph_h - 0.6, it['text'], size=c.get('text_size', 10.5),
                            line_sp=1.15, space_after=0))
        out.append(grp)
    K.anims(s, hd, 'fade', 0.0, 0.05, 0.4)
    K.anims(s, out, 'rise', 0.3, 0.12 if n > 4 else 0.18, 0.5)
    return s


# ================================================================== 19. partners: grouped logo wall
def partners(prs, ctx, c):
    """Logo wall in labelled rows. groups = [{label, logos: [photo keys]}]."""
    s = _new(prs, ctx)
    bg_light(s, ctx)
    hd = header(s, ctx, c['title'], c.get('sub', ''), c.get('tag', ''), c.get('num'))
    groups = c['groups']
    lh, gap = 0.34, 0.22
    avail = SH - 0.45 - 1.95
    rh = min(1.1, (avail - len(groups) * (lh + 0.1) - gap * (len(groups) - 1)) / len(groups))
    y = 1.95
    out = []
    for g in groups:
        lab = g['label'].upper()
        tw = K.text_w(lab, F('semi'), 10) + len(lab) * 1.2 / 72 + 0.1
        row = [K.text(s, ML, y, tw, lh, lab, size=10, font=F('semi'), color=T['primary'], anchor='m', spacing=1.2,
                      wrap=False),
               K.line(s, ML + tw + 0.15, y + lh / 2, SW - ML, y + lh / 2, T['line'], 1.0)]
        logos = g['logos']
        n = len(logos)
        lw = min(1.6, (SW - 2 * ML - 0.18 * (n - 1)) / n)
        x0 = (SW - (n * lw + (n - 1) * 0.18)) / 2
        yy = y + lh + 0.1
        for k, key in enumerate(logos):
            x = x0 + k * (lw + 0.18)
            cd = card(s, x, yy, lw, rh, radius=0.1)
            lp, _ = _contain(s, pic(ctx, key), x + 0.12, yy + 0.1, lw - 0.24, rh - 0.2)
            row.append([cd, lp])
        out.append(row)
        y = yy + rh + gap
    K.anims(s, hd, 'fade', 0.0, 0.05, 0.4)
    t = 0.3
    for row in out:
        K.anims(s, row[:2], 'wipe_l', t, 0.0, 0.4)
        K.anims(s, row[2:], 'zoom', t + 0.15, 0.06, 0.35)
        t += 0.45
    return s


# ================================================================== 20. offices: map + office list
def offices(prs, ctx, c):
    """Map image on the left with numbered pins (fx, fy as fractions of the map image), office cards on the right.
    offices = [{name, address, phone?, pin: [fx, fy]}]."""
    s = _new(prs, ctx)
    bg_light(s, ctx)
    hd = header(s, ctx, c['title'], c.get('sub', ''), c.get('tag', ''), c.get('num'))
    mp, (mx, my, mw, mh) = _contain(s, pic(ctx, c['map']), ML, 1.8, 5.9, SH - 1.8 - 0.3, nm='!!map')
    offs = c['offices'][:4]
    pins, rows = [], []
    gx = 7.0
    gw = SW - ML - gx
    n = len(offs)
    gap = 0.2
    ch = min(1.5, (SH - 0.45 - 1.95 - gap * (n - 1)) / n)
    y0 = 1.95 + ((SH - 0.45 - 1.95) - (n * ch + gap * (n - 1))) / 2
    for i, o in enumerate(offs):
        fx, fy = o['pin']
        px, py = mx + fx * mw, my + fy * mh
        gl = K.glow_img(s, px, py, 0.9, T['light'], 60)
        dot = K.oval(s, px - 0.17, py - 0.17, 0.34, fill=T['primary'], line='FFFFFF', lw=1.5)
        K.shape_text(dot, str(i + 1), size=9, color='FFFFFF', font=F('bold'), bold=B('bold'), margin=(0, 0))
        pins.append([gl, dot])
        y = y0 + i * (ch + gap)
        cd = card(s, gx, y, gw, ch)
        nb = K.oval(s, gx + 0.25, y + 0.22, 0.42, fill=T['primary'])
        K.shape_text(nb, str(i + 1), size=12, color='FFFFFF', font=F('bold'), bold=B('bold'), margin=(0, 0))
        t1 = K.text(s, gx + 0.85, y + 0.14, gw - 1.0, 0.42, o['name'], size=14, font=F('bold'), bold=B('bold'),
                    color=T['ink'], anchor='m')
        t2 = K.text(s, gx + 0.85, y + 0.56, gw - 1.0, ch - 0.6, o['address'] + (('\n' + o['phone']) if o.get('phone')
                                                                                 else ''),
                    size=10.5, font=F('body'), color=T['text'], line_sp=1.15)
        rows.append([cd, nb, t1, t2])
    K.anims(s, hd, 'fade', 0.0, 0.05, 0.4)
    K.anim(s, mp, 'fade', 0.2, 0.7)
    K.anims(s, pins, 'zoom', 0.7, 0.2, 0.4)
    K.anims(s, rows, 'fly_r', 0.8, 0.2, 0.45)
    return s


# ================================================================== credits (hidden)
def credits(prs, ctx, c):
    s = _new(prs, ctx)
    s._element.set('show', '0')
    bg_light(s, ctx)
    title_text(s, ML, 0.4, 11, 0.6, 'NGUỒN HÌNH ẢNH', 24, color=T['ink'])
    K.text(s, ML, 1.05, 12, 0.5, 'Slide ẩn — không hiển thị khi trình chiếu. ' + c['intro'], size=10, font=F('body'),
           color=T['muted'])
    lines = c['lines']
    half = (len(lines) + 1) // 2
    for col, chunk in enumerate((lines[:half], lines[half:])):
        K.text(s, ML + col * 6.1, 1.75, 5.9, 5.4, '\n'.join(chunk), size=8, font=F('body'), color=T['text'],
               line_sp=1.05)
    return s
