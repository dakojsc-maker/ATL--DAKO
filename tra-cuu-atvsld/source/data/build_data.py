#!/usr/bin/env python3
import json, re, sys
sys.path.insert(0, '.')
import mapcl
from grp_parse import parse, plain

SRC = 'data.json'
OUT = 'app_data.json'
d = json.load(open(SRC, encoding='utf-8'))

# ---- manual cluster overrides for rows where text matching is ambiguous ----
# (industry no, position prefix) -> list of cluster numbers ([] = no cluster card)
OVR = {
    (1, 'Thành viên đội sơ cấp cứu công trường'): [],
    (1, 'Khách vào công trường'): [],
    (4, 'Thành viên đội sơ cấp cứu của xưởng'): [],
    (5, 'Bảo vệ nhà máy, bảo vệ cổng'): [8],
    (5, 'Nhân viên thống kê, kiểm kê vật tư'): [9],
    (5, 'Công nhân thêu'): [11],
    (5, 'Công nhân đứng chuyền lắp ráp'): [13],
    (5, 'Công nhân vận hành máy sấy công nghiệp'): [14],
    (5, 'Nhân viên thu gom, lưu giữ chất thải nguy hại'): [17],
    (5, 'Chuyên gia, cán bộ kỹ thuật, quản lý người nước ngoài'): [18],
    (5, 'Khách vào nhà máy'): [],
    (6, 'Thành viên đội sơ cấp cứu'): [],
    (10, 'Thành viên đội sơ cấp cứu'): [],
    (10, 'Cán bộ điều tra, quy hoạch rừng'): [7],
    (11, 'Kỹ thuật viên lò sấy gỗ'): [17],
    (11, 'Bảo vệ nhà máy, bảo vệ cổng'): [8],
    (11, 'Nhân viên thống kê, kiểm kê vật tư'): [9],
    (11, 'Thợ chà nhám bằng tay'): [12],
    (11, 'Khách vào nhà máy'): [],
    (15, 'Thành viên đội sơ cấp cứu'): [],
    (15, 'Áp tải hàng giá trị'): [6],
    (15, 'Xe nâng tay không động cơ'): [8],
    (15, 'Vận hành, bảo dưỡng hệ thống lạnh'): [11],
    (16, 'Giáo viên dạy thực hành nấu ăn'): [5],
    (17, 'Lực lượng cấp cứu, cứu hộ chuyên nghiệp'): [],
}

ACR = ['HĐQT', 'HĐTV', 'ATVSLĐ', 'ATLĐ', 'PCCC', 'HSE', 'EHS', 'QA/QC', 'QA', 'QC', 'CNTT', 'R&D', 'IT', 'ICD', 'KCS', 'BIM', 'MEP',
       'QS', 'CSR', 'ISO', 'HACCP', 'GMP', 'FDI', 'CNC', 'PLC', 'BPO', 'KSNK', 'TCVN', 'NVKD', 'II', 'III', 'IV']

def sentence_case(s):
    s = s.strip()
    letters = [c for c in s if c.isalpha()]
    if not letters or not all(c == c.upper() for c in letters):
        return s
    low = s.lower()
    low = low[:1].upper() + low[1:]
    # restore acronyms
    for a in sorted(ACR, key=len, reverse=True):
        low = re.sub(r'(?<![\w])' + re.escape(a.lower()) + r'(?![\w])', a, low, flags=re.I)
    # capital after "Khối X."
    low = re.sub(r'^(Khối [A-Z]\.\s*)(\w)', lambda m: m.group(1) + m.group(2).upper(), low)
    return low

def short_title(name):
    n = re.sub(r'^NGÀNH\s+', '', name.strip(), flags=re.I)
    return n[:1].upper() + n[1:].lower() if n.isupper() else n

# ---- canonical 6 groups (from file 01, tables 2.3 and 2.4) ----
import docx
from docx.oxml.ns import qn
doc = docx.Document('../handbooks/01.docx')
tabs = doc.tables
def find_table(hdr):
    for t in tabs:
        if t.rows and [c.text.strip() for c in t.rows[0].cells][:3] == hdr:
            return t
t23 = find_table(['Nhóm', 'Đối tượng', 'Huấn luyện lần đầu (tối thiểu)'])
t24 = find_table(['Nhóm', 'Nội dung khung', 'Thông điệp giá trị (theo hồ sơ năng lực)'])
GROUPS = []
for r23, r24 in zip(t23.rows[1:], t24.rows[1:]):
    c23 = [c.text.strip() for c in r23.cells]
    c24 = r24.cells
    frame = [p.text.strip() for p in c24[1].paragraphs if p.text.strip()]
    msg = c24[2].text.strip()
    n = int(re.search(r'\d', c23[0]).group(0))
    who = c23[1]
    short = who.split(':')[0].strip()
    GROUPS.append({'n': n, 'short': short, 'who': who, 'first': c23[2], 'again': c23[3], 'frame': frame, 'msg': msg})

# ---- industries ----
INDS = []
typed_hdr = set()
for ind in d:
    use_ref, low = mapcl.assign(ind)
    rows_total = 0
    blocks_out = []
    for bi, b in enumerate(ind['blocks']):
        bname = b['name']
        m = re.match(r'^(?:KHỐI|Khối)\s+([A-Z])\s*[\.:]\s*(.*)$', bname)
        bid = m.group(1) if m else chr(65 + bi)
        btitle = sentence_case(m.group(2)) if m else sentence_case(bname)
        rows_out = []
        for ri, r in enumerate(b['rows']):
            gp = parse(r['grp'], r['note'], r.get('dia'))
            cl = list(r['cl'])
            # overrides and rules
            for (no, pref), val in OVR.items():
                if no == ind['no'] and r['pos'].startswith(pref):
                    cl = list(val)
            gtxt = plain(r['grp']).lower()
            if gtxt.startswith('không thuộc') or 'không thuộc 6 nhóm' in gtxt[:30]:
                cl = []
            if r.get('src') == 'match':
                cl = cl[:1]
            if 'alt' not in gp['f'] and re.search(r'hoặc|và/hoặc', plain(r['grp'])) and len(gp['g']) + len(gp['o']) > 1:
                gp['f'].append('alt')
            # classify optionals
            opts = []
            for o in gp['o']:
                w = o['w']
                kind = 'rec' if re.search(r'khuyến nghị|thường thêm|thường\s+Nhóm', w, re.I) else 'cond'
                opts.append({'g': o['g'], 'w': w[:260], 'k': kind})
            rows_out.append({
                'id': f"{ind['no']:02d}.{bid}{ri + 1}",
                'p': r['pos'],
                'raw': plain(r['grp']),
                'g': gp['g'],
                'o': opts,
                'x': gp['x'],
                'f': gp['f'],
                'n': r['note'],
                'cl': cl,
            })
            rows_total += 1
        blocks_out.append({'id': bid, 'name': btitle, 'rows': rows_out})
    # types
    types = None
    if ind['types']:
        h = ind['types']['header']
        typed_hdr.add(tuple(h))
        types = []
        for r in ind['types']['rows']:
            r = r + [''] * (4 - len(r))
            types.append({'n': r[0], 'd': r[1], 'need': r[2], 'who': r[3]})
    # scale tables
    sc = []
    for t in ind['scale']['tables']:
        sc.append({'t': t['title'], 'h': t['head'], 'ranged': t['ranged'],
                   'r': [{'l': x['label'], 'c': x['cells'], 'lo': x['lo'], 'hi': x['hi']} for x in t['rows']]})
    clusters = []
    for c in ind['clusters']:
        clusters.append({
            'no': c['no'], 't': c['title'], 'pos': c['positions'], 'grp': c['groups'], 'con': c['content'],
            'faq': c['faq'], 'soft': c['soft'], 'up': c['upsell'], 'note': c['note'],
        })
    INDS.append({
        'id': f"{ind['no']:02d}", 'no': ind['no'], 'name': short_title(ind['name']), 'full': ind['name'], 'legal': ind['legal'],
        'types': types, 'scale': {'t': ind['scale']['title'], 'i': ind['scale']['intro'], 'tables': sc},
        'blocks': blocks_out, 'clusters': clusters,
    })


# ---- làm sạch lời giải thích nhóm có điều kiện (hiển thị cho NVKD) ----
def _nums(t):
    out = set()
    for m in re.finditer(r'Nhóm\s*((?:\d(?:\s*(?:,|và|hoặc|/|và/hoặc|đến|-|–)\s*)?)+)', t):
        for x in re.findall(r'\d', m.group(1)): out.add(int(x))
    return out
W_FIX = {
    ('02.L1', 4): ('cond', 'Nếu là kỹ thuật viên thì Nhóm 4'),
    ('14.B2', 1): ('cond', 'Xem xét Nhóm 1 nếu chỉ huy cả ca'),
    ('14.J2', 3): ('cond', 'Nhóm 3 nếu là người vận hành xe nâng (xe nâng thuộc Danh mục nghiêm ngặt)'),
    ('17.B1', 5): ('cond', 'Có thể thêm Nhóm 5 tùy công việc chuyên môn'),
    ('13.O4', 3): ('cond', 'Nhóm 3 nếu làm việc tại mỏ, nhà máy; không thì Nhóm 4'),
    ('11.B4', 4): ('alt', 'Tùy công việc thực tế của người đó'),
    ('11.B4', 6): ('rec', 'Khuyến nghị cử làm an toàn, vệ sinh viên (thêm Nhóm 6)'),
    ('05.B4', 6): ('rec', 'Khuyến nghị cử làm an toàn, vệ sinh viên (thêm Nhóm 6)'),
    ('16.B1', 6): ('rec', 'Khuyến nghị cử làm an toàn, vệ sinh viên (thêm Nhóm 6)'),
    ('14.B1', 3): ('rec', 'Thường thêm Nhóm 3'),
    ('17.A1', 3): ('cond', 'Chủ phòng khám trực tiếp khám bệnh thì thêm Nhóm 3 và/hoặc Nhóm 5 theo công việc'),
    ('17.A1', 5): ('cond', 'Chủ phòng khám trực tiếp khám bệnh thì thêm Nhóm 3 và/hoặc Nhóm 5 theo công việc'),
    ('07.B3', 4): ('alt', 'Nếu không chỉ huy một bộ phận thì Nhóm 3 hoặc Nhóm 4 theo công việc'),
    ('03.D3', 4): ('alt', 'Nhóm 2 nếu có giám sát ATVSLĐ, còn lại Nhóm 4'),
    ('13.A4', 2): ('alt', 'Nhóm 1 hoặc Nhóm 2 theo nhiệm vụ'),
    ('08.B2', 3): ('cond', 'Thêm Nhóm 3 nếu trực tiếp làm công việc nghiêm ngặt (ví dụ tổ trưởng kỹ thuật)'),
    ('08.E1', 4): ('alt', 'Nhóm 3 có khả năng cao hơn, một số trường hợp Nhóm 4; chuyên viên xác nhận theo máy thực tế'),
    ('01.I1', 4): ('alt', 'Tùy công việc thực tế của người đó'),
}
def _clean(w):
    w = re.sub(r'\s+', ' ', w).strip().rstrip(';,. ').strip()
    return w[:1].upper() + w[1:] if w else w
for _ind in INDS:
    for _b in _ind['blocks']:
        for _r in _b['rows']:
            _m = _nums(_r['raw'])
            for _o in _r['o']:
                _o['w'] = _clean(_o['w'])
                key = (_r['id'], _o['g'])
                if key in W_FIX:
                    _o['k'], _o['w'] = W_FIX[key]
                elif 'alt' in _r['f'] and _o['g'] in _m and 'hoặc' in _o['w'].lower() and not re.search(r'\b(nếu|khi)\b', _o['w'].lower()):
                    if 'và/hoặc' in _r['raw']:
                        _o['k'] = 'cond'
                        _o['w'] = 'Nhóm ' + ' và/hoặc Nhóm '.join(str(x) for x in sorted(set(_r['g']) | {_o['g']})) + ' tùy công việc thực tế'
                    else:
                        _o['k'] = 'alt'
                        _o['w'] = 'Tùy công việc thực tế của người đó'


# ---- sửa văn bản còn sót: tham chiếu "nhóm N" của Danh mục -> "mục N", câu bị cắt, tham chiếu khối ----
W_FIX2 = {
    ('02.C4', 3): 'Xem xét Nhóm 3 nếu làm công việc thuộc Danh mục (mục 22, 23)',
    ('02.L5', 3): 'Xem xét Nhóm 3 nếu làm công việc thuộc Danh mục (mục 22, 23)',
    ('02.C6', 3): 'Xem xét Nhóm 3 nếu làm công việc thuộc Danh mục (mục 23)',
    ('02.J1', 3): 'Xem xét Nhóm 3 nếu làm công việc thuộc Danh mục (mục 13; dùng hóa chất nguy hiểm: mục 2)',
    ('07.M1', 3): 'Xem xét Nhóm 3 nếu làm công việc thuộc Danh mục (mục 13; dùng hóa chất nguy hiểm: mục 2)',
    ('13.N1', 3): 'Xem xét Nhóm 3 nếu làm công việc thuộc Danh mục (mục 13; dùng hóa chất nguy hiểm: mục 2)',
    ('02.K2', 3): 'Xem xét Nhóm 3 tùy công việc thực tế',
    ('02.F3', 3): 'Nhóm 3 nếu bốc xếp từ 30 kg (mục 20)',
    ('07.F3', 3): 'Nhóm 3 nếu bốc xếp từ 30 kg (mục 20)',
    ('13.F4', 3): 'Nhóm 3 nếu bốc xếp từ 30 kg (mục 20)',
    ('08.G1', 3): 'Thường thêm Nhóm 3 nếu trực tiếp làm công việc thuộc Danh mục',
    ('14.B1', 3): 'Thường thêm Nhóm 3 nếu trực tiếp làm công việc nghiêm ngặt',
    ('10.B3', 3): 'Nhóm 3 nếu trực tiếp làm công việc hóa chất nguy hiểm (mục 2) hoặc thuộc mục 22, 23 (xét nghiệm vi sinh vật; kiểm nghiệm, giữ giống, ký sinh trùng)',
}
for _ind in INDS:
    for _b in _ind['blocks']:
        for _r in _b['rows']:
            for _o in _r['o']:
                key = (_r['id'], _o['g'])
                if key in W_FIX2:
                    _o['w'] = W_FIX2[key]
                    if key == ('14.B1', 3): _o['k'] = 'cond'
                w = _o['w']
                w = re.sub(r'\s*\(xem Khối [A-Z]\)', '', w)
                w = re.sub(r'\(nhóm (\d+)\)', r'(mục \1)', w)
                w = re.sub(r'thuộc nhóm (\d+(?:, \d+)*)', r'thuộc mục \1', w)
                w = re.sub(r'\s+', ' ', w).strip()
                _o['w'] = w[:1].upper() + w[1:]

print('types headers:', typed_hdr)
OUTD = {'groups': GROUPS, 'inds': INDS}
s = json.dumps(OUTD, ensure_ascii=False, separators=(',', ':'))
open(OUT, 'w', encoding='utf-8').write(s)
print('bytes', len(s.encode('utf-8')))
for i in INDS:
    nrows = sum(len(b['rows']) for b in i['blocks'])
    nocl = sum(1 for b in i['blocks'] for r in b['rows'] if not r['cl'])
    nodef = sum(1 for b in i['blocks'] for r in b['rows'] if not r['g'])
    print(i['id'], i['name'], 'rows', nrows, 'no-cluster', nocl, 'no-default-group', nodef, 'blocks', [b['id'] for b in i['blocks']][:3], '...', 'scale', [(t['ranged']) for t in i['scale']['tables']])
