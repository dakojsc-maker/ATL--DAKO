import base64, io, json
import docx
from playwright.sync_api import sync_playwright
import os
HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(os.path.dirname(HERE))          # thư mục gói
OUTD = os.path.join(HERE, 'out'); os.makedirs(OUTD, exist_ok=True)   # ảnh chụp, file docx thử
CHROME = '/opt/pw-browsers/chromium' if os.path.exists('/opt/pw-browsers/chromium') else None

# Kiểm thử bản artifact (python3 build.py --artifact): bọc khung trang giống claude.ai,
# giả lập window.claude.use("downloads") rồi chạy luồng xuất Word, kiểm tra nút In / Chia sẻ bị ẩn.
SRC = os.path.join(ROOT, 'artifact', 'tra-cuu-atvsld.html')
page = open(SRC, encoding='utf-8').read()
assert not any(x in page[:20000].lower() for x in ('<!doctype', '<html', '<head>', '<body')), 'bản artifact không được có khung trang riêng'
assert page.startswith('<title>'), 'thẻ <title> phải đứng đầu file'
WRAP = os.path.join(OUTD, 'artifact_wrapped.html')
open(WRAP, 'w', encoding='utf-8').write(
    '<!doctype html><html><head><meta charset=utf8><meta name=viewport content="width=device-width,initial-scale=1,viewport-fit=cover">'
    '<style>:root{color-scheme:light}body{margin:0;font:14px system-ui}img{max-width:100%}[hidden]{display:none!important}</style></head><body>'
    + page + '</body></html>')
URL = 'file://' + WRAP

# mode: ok (người xem đồng ý), declined (từ chối), none (không có quyền downloads)
MOCK = """(mode)=>{ window.__saves=[]; window.claude={ use:(name)=>new Promise(res=>setTimeout(()=>{
  if(name!=='downloads'||mode==='none') return res(null);
  res(Object.freeze({ save:async(req)=>{ const buf=new Uint8Array(await req.data.arrayBuffer()); let s='';
    for(let i=0;i<buf.length;i+=8192) s+=String.fromCharCode.apply(null,buf.subarray(i,i+8192));
    window.__saves.push({filename:req.filename,b64:btoa(s)});
    if(mode==='declined') throw {code:'declined',message:'no'}; return {status:'saved'}; } }));
},30)) }; }"""

errs, bad = [], []
def clickjs(pg, sel, i=0):
    pg.evaluate("([s,i])=>{const e=document.querySelectorAll(s)[i]; e.scrollIntoView({block:'center'}); e.click()}", [sel, i])
def flow(b, mode, vp, mobile):
    ctx = b.new_context(viewport=vp, device_scale_factor=2 if mobile else 1, is_mobile=mobile, has_touch=mobile)
    ctx.add_init_script(script='(' + MOCK + ')(' + json.dumps(mode) + ')')
    pg = ctx.new_page()
    pg.on('pageerror', lambda e: errs.append(str(e)))
    pg.on('console', lambda m: errs.append(m.text) if m.type == 'error' else None)
    pg.goto(URL); pg.wait_for_timeout(300)
    clickjs(pg, 'button[data-act="ind"][data-id="01"]'); pg.wait_for_timeout(150)
    pg.fill('#f-q', 'công nhân'); pg.wait_for_timeout(200)
    n = pg.inner_text('#find-top .cnt')
    if '21 vị trí' not in n: bad.append([mode, 'search', n])
    for i in [0, 1, 2]: clickjs(pg, '#find-list .add', i)
    pg.fill('#f-q', 'phụ hồ'); pg.wait_for_timeout(150); clickjs(pg, '#find-list .add', 0)
    clickjs(pg, '.bottom-nav button[data-view="list"]' if mobile else '#tabs button[data-view="list"]'); pg.wait_for_timeout(300)
    acts = pg.eval_on_selector_all('#list-top .res-actions button', 'e=>e.map(x=>x.dataset.act)')
    if 'print' in acts: bad.append([mode, 'print button shown'])
    if 'export' not in acts or 'copy' not in acts: bad.append([mode, 'missing actions', acts])
    if mobile: pg.screenshot(path=OUTD + '/art_mb_list.png')
    clickjs(pg, 'button[data-act="export"]'); pg.wait_for_timeout(200)
    if pg.locator('button[data-act="x-share"]').count(): bad.append([mode, 'share button shown'])
    clickjs(pg, '#x-go'); pg.wait_for_timeout(1500)
    saves = pg.evaluate('window.__saves')
    toast = pg.inner_text('#toast')
    sheet_open = pg.evaluate("document.getElementById('sheet').classList.contains('open')")
    print(mode, 'mobile' if mobile else 'desktop', '| saves', [s['filename'] for s in saves], '| toast:', toast, '| sheet open:', sheet_open,
          '| overflow', pg.evaluate("document.documentElement.scrollWidth-document.documentElement.clientWidth"))
    if sheet_open: bad.append([mode, 'sheet still open'])
    if mode == 'none':
        if saves or 'chưa được phép' not in toast: bad.append([mode, 'none path', toast])
    else:
        if len(saves) != 1 or not saves[0]['filename'].endswith('.docx'): bad.append([mode, 'save not called', saves])
        else:
            d = docx.Document(io.BytesIO(base64.b64decode(saves[0]['b64'])))
            txt = '\n'.join(p.text for p in d.paragraphs)
            if 'ĐỀ XUẤT' not in '\n'.join(c.text for t in d.tables for r in t.rows for c in r.cells) or 'II. BẢNG NHÂN SỰ' not in txt:
                bad.append([mode, 'docx content'])
            if mode == 'ok':
                open(OUTD + '/art_out.docx', 'wb').write(base64.b64decode(saves[0]['b64']))
                if not toast.startswith('Đã lưu file'): bad.append([mode, 'toast', toast])
            if mode == 'declined' and toast != 'Đã hủy lưu file.': bad.append([mode, 'toast', toast])
    ctx.close()

with sync_playwright() as p:
    b = p.chromium.launch(executable_path=CHROME)
    flow(b, 'ok', {'width': 1366, 'height': 900}, False)
    flow(b, 'declined', {'width': 1366, 'height': 900}, False)
    flow(b, 'none', {'width': 1366, 'height': 900}, False)
    flow(b, 'ok', {'width': 390, 'height': 844}, True)
    b.close()
print('BAD', bad)
print('ERRS', errs)
