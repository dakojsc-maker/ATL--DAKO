import base64, io
import docx
from playwright.sync_api import sync_playwright
import os
HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(os.path.dirname(HERE))          # thư mục gói
OUTD = os.path.join(HERE, 'out'); os.makedirs(OUTD, exist_ok=True)   # ảnh chụp, file docx thử
URL = 'file://' + ROOT + '/Tra_cuu_nhom_huan_luyen_ATVSLD_VienSTP_DAKO.html'
CHROME = '/opt/pw-browsers/chromium' if os.path.exists('/opt/pw-browsers/chromium') else None

# Kiểm thử chống lặp ý và các đầu ra phụ:
#  1. vị trí không có nhóm mặc định: không có hai dòng cùng ý trên màn hình tra cứu, thẻ danh sách không bảo "chọn nhóm" với người ngoài 6 nhóm
#  2. câu giải thích cạnh nhãn nhóm không bắt đầu lặp lại đúng nhãn đó ("Nhóm 3 Nhóm 3 nếu…")
#  3. tóm tắt sao chép có "hoặc Nhóm X" như file Word
#  4. in: mọi mục nội dung được mở khi in và đóng lại sau khi in
#  5. chỉ chọn vị trí chưa có nhóm: màn hình và file Word không có bảng rỗng
errs, bad = [], []
with sync_playwright() as p:
    b = p.chromium.launch(executable_path=CHROME)
    pg = b.new_context(viewport={'width': 1366, 'height': 900}).new_page()
    pg.on('pageerror', lambda e: errs.append(str(e)))
    pg.on('console', lambda m: errs.append(m.text) if m.type == 'error' else None)
    pg.goto(URL); pg.wait_for_timeout(300)
    res = pg.evaluate("""()=>{
      const bad=[]; let nNo=0, nOpt=0;
      const nvf=s=>s.normalize('NFD').replace(/[\\u0300-\\u036f]/g,'').replace(/đ/g,'d').toLowerCase();
      for(const ind of STP.IND){
        STP.S.ind=ind.id; STP.S.picker=false; STP.S.q=''; STP.S.selBy[ind.id]={}; STP.renderAll();
        document.querySelectorAll('#find-list .rr').forEach(el=>{
          const r=STP.ROW[el.dataset.id];
          const lines=[...el.querySelectorAll('.rr-groups .gl')].map(g=>g.innerText.replace(/\\s+/g,' ').trim());
          if(!r.g.length){
            nNo++;
            const flags=lines.filter(t=>!/^\\+ |^hoặc /.test(t));
            // hai dòng mô tả không được nói cùng một ý
            for(let i=0;i<flags.length;i++) for(let j=i+1;j<flags.length;j++){
              const a=nvf(flags[i]), c=nvf(flags[j]);
              if(a.includes('theo cong viec')&&c.includes('theo cong viec')) bad.push([r.id,'dup work',flags]);
              if(a.includes('ranh gioi')&&c.includes('ranh gioi')) bad.push([r.id,'dup rev',flags]);
              if(a.includes('chuyen vien')&&c.includes('chuyen vien')) bad.push([r.id,'dup cv',flags]);
            }
          }
          el.querySelectorAll('.gl.cd').forEach(g=>{
            nOpt++;
            const badges=[...g.querySelectorAll('.gb')].map(x=>x.textContent.trim());
            const why=g.querySelector('span:last-child').textContent.trim();
            if(badges.length===1 && new RegExp('^(Thêm )?'+badges[0]+' (nếu|khi)').test(why)) bad.push([r.id,'stutter',why]);
          });
        });
        // thẻ danh sách
        ind.blocks.forEach(b=>b.rows.forEach(r=>{ if(!r.g.length) STP.pickRow(r.id,true); }));
        STP.S.view='list'; STP.renderAll();
        document.querySelectorAll('#list-rows .sc').forEach(sc=>{
          const r=STP.ROW[sc.dataset.id]; const t=(sc.querySelector('.rawtxt')||{}).textContent||'';
          const outside=r.f.includes('none')||r.f.includes('svc');
          if(outside && /Bật nhóm|Chọn nhóm/.test(t)) bad.push([r.id,'card outside asks to pick group',t]);
          if(/theo công việc thực tế.*theo công việc thực tế/.test(t)) bad.push([r.id,'card dup',t]);
          sc.querySelectorAll('.cond').forEach(c=>{
            const badges=[...c.querySelectorAll('.gb')].map(x=>x.textContent.trim()), why=c.querySelector('.why').textContent.trim();
            if(badges.length===1 && new RegExp('^(Thêm )?'+badges[0]+' (nếu|khi)').test(why)) bad.push([r.id,'card stutter',why]);
          });
        });
        // chỉ có vị trí chưa có nhóm: không có bảng tổng hợp rỗng
        if(!STP.buildModel().usedGroups.length && document.querySelector('#list-out table')) bad.push([ind.id,'empty summary table']);
        STP.S.selBy[ind.id]={};
      }
      return {nNo,nOpt,bad};
    }""")
    print('no-group rows', res['nNo'], 'option lines', res['nOpt'], 'bad', len(res['bad']), res['bad'][:10])
    bad += res['bad']
    # tóm tắt: có "hoặc Nhóm 4" cho phụ hồ
    s = pg.evaluate("""()=>{STP.S.ind='01'; STP.S.selBy['01']={}; STP.pickRow('01.I1',true); STP.renderAll(); return STP.summaryText(STP.buildModel());}""")
    line = [l for l in s.split('\n') if l.startswith('1. ')][0]
    print('summary:', line)
    if 'hoặc Nhóm 4' not in line: bad.append('summary alt missing')
    # in: mở hết, in xong đóng lại
    pg.evaluate("document.querySelector('#tabs button[data-view=list]').click()"); pg.wait_for_timeout(200)
    st = pg.evaluate("""()=>{const c=()=>[...document.querySelectorAll('details.cl')].filter(d=>!d.open).length;
      const a=c(); window.dispatchEvent(new Event('beforeprint')); window.dispatchEvent(new Event('beforeprint'));
      const m=c(); window.dispatchEvent(new Event('afterprint')); return [a,m,c()];}""")
    print('details closed before/during/after print:', st)
    if not (st[0] > 0 and st[1] == 0 and st[2] == st[0]): bad.append(['print details', st])
    # Word khi chỉ có vị trí chưa có nhóm
    b64 = pg.evaluate("""async()=>{STP.S.ind='07'; STP.S.selBy['07']={}; ['07.B4','07.O5'].forEach(id=>STP.pickRow(id,true)); STP.renderAll();
      const blob=await STP.makeDoc(STP.presetOpts('cust')); const buf=new Uint8Array(await blob.arrayBuffer()); let s='';
      for(let i=0;i<buf.length;i+=8192) s+=String.fromCharCode.apply(null,buf.subarray(i,i+8192)); return btoa(s);}""")
    raw = base64.b64decode(b64); open(OUTD + '/nogroup_07.docx', 'wb').write(raw)
    d = docx.Document(io.BytesIO(raw))
    heads = [t for t in d.tables if t.rows and t.rows[0].cells[0].text.strip() == 'Nhóm']
    texts = [p.text for p in d.paragraphs]
    print('word: group tables', len(heads), '| has IV:', any(t.startswith('IV.') for t in texts))
    if heads: bad.append('word empty group table')
    b.close()
print('BAD', bad)
print('ERRS', errs)
